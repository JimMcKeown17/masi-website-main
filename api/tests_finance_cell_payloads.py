"""Payload parity at the authenticated raw upload boundary, for both kinds."""
from datetime import datetime
from io import BytesIO
from urllib.parse import urlencode
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from api.finance_budget_test_utils import budget_workbook, budget_ledger
from api.finance_run_test_utils import actor
from api.models import FinanceRun
from api.parsers import finance_workbook as parser
from api.services import finance_runs
from api.tests_finance_upload_safety import NAME, rewrite, workbook_bytes


class CellPayloadUploadTests(TestCase):
    def setUp(self):
        self.user = actor()
        self.dependency = budget_ledger(self.user)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def upload(self, kind, data):
        query = dict(kind=kind, year=2026, source_name=NAME)
        if kind == 'budgets':
            query['ledger_run_id'] = str(self.dependency.pk)
        return self.client.post('/api/finance/runs/?' + urlencode(query), data,
                                content_type=parser.MIME)

    def fixture(self, kind, payload, shared=False, header=False, coordinate_override=None):
        data = budget_workbook() if kind == 'budgets' else workbook_bytes()
        coordinate = ('A3' if kind == 'budgets' else 'A1') if header else ('L1' if kind == 'budgets' else 'A2')
        coordinate = coordinate_override or coordinate
        def change(xml):
            root = ET.fromstring(xml)
            cell = root.find('.//' + parser.NS + 'c' + f"[@r='{coordinate}']")
            for child in list(cell):
                cell.remove(child)
            if shared:
                cell.set('t', 's')
                ET.SubElement(cell, parser.NS + 'v').text = '0'
            else:
                content = ET.fromstring(f'<c xmlns="{parser.NS[1:-1]}">{payload}</c>')
                cell.set('t', 'inlineStr' if content.find(parser.NS + 'is') is not None else 'n')
                cell.extend(content)
            return ET.tostring(root)
        changes = {'xl/worksheets/sheet1.xml': change}
        extra = []
        if shared:
            changes['[Content_Types].xml'] = lambda x: x.replace(b'</Types>', b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            extra = [('xl/sharedStrings.xml', f'<sst xmlns="{parser.NS[1:-1]}"><si>{payload}</si></sst>'.encode())]
        return rewrite(data, changes, extra), coordinate

    def assert_refused(self, kind, payload, shared=False):
        data, _ = self.fixture(kind, payload, shared)
        before = set(FinanceRun.objects.values_list('pk', flat=True))
        name = 'build_budget_run_artifact' if kind == 'budgets' else 'build_run_artifact'
        with patch.object(finance_runs, name, wraps=getattr(finance_runs, name)) as producer:
            response = self.upload(kind, data)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(response.data['code'], 'XML_INVALID')
        producer.assert_not_called()
        self.assertEqual(set(FinanceRun.objects.values_list('pk', flat=True)), before)

    def assert_control(self, kind, payload, expected, shared=False, header=False, formula=None, coordinate_override=None):
        data, coordinate = self.fixture(kind, payload, shared, header, coordinate_override)
        name = 'build_budget_run_artifact' if kind == 'budgets' else 'build_run_artifact'
        original = getattr(finance_runs, name)
        def inspect(stream, *args, **kwargs):
            position = stream.tell()
            for data_only in (True, False):
                wb = load_workbook(stream, read_only=True, data_only=data_only)
                try:
                    self.assertEqual(wb.worksheets[0][coordinate].value,
                                     formula if formula and not data_only else expected)
                finally:
                    wb.close()
            stream.seek(position)
            return original(stream, *args, **kwargs)
        with patch.object(finance_runs, name, side_effect=inspect) as producer:
            response = self.upload(kind, data)
        self.assertEqual(response.status_code, 201, response.data)
        producer.assert_called_once()
        run = FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status, 'candidate', run.failure)
        return data

    def test_rich_header_scanner_matches_openpyxl(self):
        # Text.content puts plain text first even when XML puts runs first;
        # phonetic text contributes no characters and shared escapes are removed.
        payload = '<r><rPr><b/></rPr><t>te</t></r><t>Dax005F_</t><rPh sb="0" eb="4"><t>ignored</t></rPh><phoneticPr fontId="0"/>'
        for shared in (False, True):
            with self.subTest(shared=shared):
                text = payload if shared else payload.replace('x005F_', '')
                data = self.assert_control('funders', text if shared else '<is>' + text + '</is>', 'Date', shared, True)
                with ZipFile(BytesIO(data)) as archive, patch.object(parser, '_label', wraps=parser._label) as label:
                    strings = parser._shared_strings(archive)
                    parser._scan_sheet(archive, 'xl/worksheets/sheet1.xml', 'Expenditure', strings)
                self.assertIn(('Date',), [call.args for call in label.call_args_list])


BAD_PAYLOADS = {
    'wrapped_v': ('<x><v>46024</v></x>', False),
    'duplicate_v': ('<v/><v>46024</v>', False),
    'nested_v': ('<v>46024<t>hidden</t></v>', False),
    'unknown_cell_child': ('<v>46024</v><x/>', False),
    'foreign_cell_child': ('<v>46024</v><x xmlns="urn:foreign"/>', False),
    'wrapped_inline_t': ('<is><x><t>hidden</t></x></is>', False),
    'unknown_inline_child': ('<is><t>value</t><x/></is>', False),
    'wrapped_shared_t': ('<x><t>hidden</t></x>', True),
    'unknown_shared_child': ('<t>value</t><x/>', True),
    'duplicate_f': ('<f>1</f><f>2</f><v>46024</v>', False),
    'nested_f': ('<f>1<x/></f><v>46024</v>', False),
    'duplicate_is': ('<is><t>a</t></is><is><t>b</t></is>', False),
    'duplicate_plain_t': ('<t>a</t><t>b</t>', True),
    'duplicate_run_t': ('<r><t>a</t><t>b</t></r>', True),
    'nested_run_properties': ('<r><rPr><b><x/></b></rPr><t>a</t></r>', True),
    'foreign_shared_child': ('<t xmlns="urn:foreign">hidden</t>', True),
}


def refusal(kind, payload, shared):
    def test(self):
        self.assert_refused(kind, payload, shared)
    return test


def controls(kind):
    def test(self):
        self.assert_control(kind, '<v>46024</v>', datetime(2026, 1, 2))
        expected = 'Category' if kind == 'budgets' else 'Date'
        rich = '<r><rPr><b/><color rgb="FF000000"/></rPr><t>' + expected[:2] + '</t></r><r><t>' + expected[2:] + '</t></r>'
        for shared in (False, True):
            for phonetic in (False, True):
                with self.subTest(shared=shared, phonetic=phonetic):
                    text = rich + ('<rPh sb="0" eb="3"><t>ignored</t></rPh><phoneticPr fontId="0"/>' if phonetic else '')
                    self.assert_control(kind, text if shared else '<is>' + text + '</is>', expected, shared, header=True)
        if kind == 'budgets':
            self.assert_control(kind, '<f>INT(MONTH(L1))</f><v>3</v>', 3,
                                formula='=INT(MONTH(L1))', coordinate_override='M1')
        else:
            self.assert_control(kind, '<f>46023+1</f><v>46024</v>', datetime(2026, 1, 2), formula='=46023+1')
    return test


for _kind in ('budgets', 'funders'):
    for _shape, (_payload, _shared) in BAD_PAYLOADS.items():
        setattr(CellPayloadUploadTests, f'test_{_kind}_{_shape}_refused', refusal(_kind, _payload, _shared))
    setattr(CellPayloadUploadTests, f'test_{_kind}_payload_controls', controls(_kind))


class CellTypePayloadUploadTests(TestCase):
    setUp = CellPayloadUploadTests.setUp
    upload = CellPayloadUploadTests.upload

    def fixture(self, kind, cell_type, payload, *, control=False):
        data = budget_workbook() if kind == 'budgets' else workbook_bytes()
        # Refusals reproduce the reviewer's ordered, otherwise empty row 8.
        # Controls occupy unused cells so business rules do not obscure parsing.
        coordinate = ('H1' if kind == 'budgets' else 'Z1') if control else 'F8'
        sheet = 2 if control and kind == 'funders' else 1
        def change(xml):
            root = ET.fromstring(xml)
            rows = root.find(parser.NS + 'sheetData')
            number = '1' if control else '8'
            row = rows.find(parser.NS + f"row[@r='{number}']")
            if row is None:
                row = ET.SubElement(rows, parser.NS + 'row', r=number)
            cell = ET.fromstring(f'<c xmlns="{parser.NS[1:-1]}" r="{coordinate}" s="0">{payload}</c>')
            if cell_type is not None:
                cell.set('t', cell_type)
            row.append(cell)
            row[:] = sorted(row, key=lambda c: parser._coordinate(c.get('r'))[1])
            root.find(parser.NS + 'dimension').set('ref', 'A1:AU8' if kind == 'budgets' else 'A1:Z8')
            return ET.tostring(root)
        changes = {f'xl/worksheets/sheet{sheet}.xml': change,
                   '[Content_Types].xml': lambda x: x.replace(b'</Types>', b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')}
        extra = [('xl/sharedStrings.xml', f'<sst xmlns="{parser.NS[1:-1]}"><si><t>Synthetic text</t></si></sst>'.encode())]
        return rewrite(data, changes, extra), sheet - 1, coordinate


TYPE_REFUSALS = {
    'inline_v': ('inlineStr', '<v>999</v>'),
    'numeric_is': ('n', '<is><t>999</t></is>'),
    'shared_is': ('s', '<is><t>999</t></is>'),
    'shared_index_outside_table': ('s', '<v>1</v>'),
    'shared_negative_index': ('s', '<v>-1</v>'),
    'shared_noninteger_index': ('s', '<v>0.5</v>'),
    'unknown_type': ('unknown', '<v>999</v>'),
    'inline_is_and_v': ('inlineStr', '<is><t>999</t></is><v>999</v>'),
    **{f'{name}_is': (kind, '<is><t>999</t></is>') for name, kind in
       [('default', None), ('boolean', 'b'), ('date', 'd'), ('error', 'e'), ('string', 'str')]},
}
TYPE_CONTROLS = {
    'numeric': ('n', '<v>999</v>', 999, None),
    'default_numeric': (None, '<v>999</v>', 999, None),
    'formula_cache': ('n', '<f>1+2</f><v>3</v>', 3, '=1+2'),
    'inline': ('inlineStr', '<is><t>Synthetic text</t></is>', 'Synthetic text', None),
    'shared': ('s', '<v>0</v>', 'Synthetic text', None),
    'boolean': ('b', '<v>1</v>', True, None),
    'date': ('d', '<v>2026-01-02T00:00:00</v>', datetime(2026, 1, 2), None),
    'error': ('e', '<v>#NAME?</v>', '#NAME?', None),
    'string': ('str', '<v>Synthetic text</v>', 'Synthetic text', None),
    'empty_styled': (None, '', None, None),
    'empty_shared': ('s', '', None, None),
}


def type_refusal(kind, cell_type, payload):
    def test(self):
        data, _, _ = self.fixture(kind, cell_type, payload)
        before = set(FinanceRun.objects.values_list('pk', flat=True))
        name = 'build_budget_run_artifact' if kind == 'budgets' else 'build_run_artifact'
        with patch.object(finance_runs, name, wraps=getattr(finance_runs, name)) as producer:
            response = self.upload(kind, data)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(response.data['code'], 'XML_INVALID')
        producer.assert_not_called()
        self.assertEqual(set(FinanceRun.objects.values_list('pk', flat=True)), before)
    return test


def type_control(kind, cell_type, payload, expected, formula):
    def test(self):
        data, sheet, coordinate = self.fixture(kind, cell_type, payload, control=True)
        name = 'build_budget_run_artifact' if kind == 'budgets' else 'build_run_artifact'
        original = getattr(finance_runs, name)
        def inspect(stream, *args, **kwargs):
            position = stream.tell()
            for data_only in (True, False):
                wb = load_workbook(stream, read_only=True, data_only=data_only)
                try:
                    actual = wb.worksheets[sheet][coordinate].value
                    wanted = formula if formula and not data_only else expected
                    self.assertEqual(actual, wanted)
                    self.assertIs(type(actual), type(wanted))
                finally:
                    wb.close()
            stream.seek(position)
            return original(stream, *args, **kwargs)
        with patch.object(finance_runs, name, side_effect=inspect) as producer:
            response = self.upload(kind, data)
        self.assertEqual(response.status_code, 201, response.data)
        producer.assert_called_once()
        run = FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status, 'candidate', run.failure)
    return test


for _kind in ('budgets', 'funders'):
    for _shape, _args in TYPE_REFUSALS.items():
        setattr(CellTypePayloadUploadTests, f'test_{_kind}_{_shape}_refused', type_refusal(_kind, *_args))
    for _shape, _args in TYPE_CONTROLS.items():
        setattr(CellTypePayloadUploadTests, f'test_{_kind}_{_shape}_control', type_control(_kind, *_args))
