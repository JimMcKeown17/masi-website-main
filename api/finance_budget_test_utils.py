"""In-memory, invented budget workbook; no operating-source labels or amounts."""
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from api.tests_finance_upload_safety import rewrite


def budget_workbook(*, sheets=3, unsized=False, half=False):
    wb = Workbook()
    sheet = wb.active
    sheet.title = '2026 Budget'
    for col, value in {1:'Category', 2:'Cat', 3:'Calc', 4:'WF', 5:'BC',
                       6:'2026 Budget', 7:'2026 Actual ', 47:'Synthetic column'}.items():
        sheet.cell(3, col, value)
    sheet['L1'] = datetime(2026, 3, 1)
    sheet['M1'] = '=INT(MONTH(L1))'
    for row, level, label in [(4,1,'Synthetic department'), (5,2,'Synthetic section'),
                              (6,3,'Synthetic line')]+([(7,3,'Synthetic second')] if half else []):
        sheet.cell(row,1,label); sheet.cell(row,2,level)
        if level < 3:
            end = 7 if half and row == 5 else row+1
            for col in ('F','G','L','M','N'):
                sheet[f'{col}{row}'] = f'=SUM({col}{row+1}:{col}{end})'
        else:
            sheet[f'E{row}'] = 'SYNTHETIC'
            sheet[f'F{row}'] = 1.2345
            sheet[f'G{row}'] = f'=IFERROR(VLOOKUP(E{row},\'Actual 2026\'!A:B,2,0),"")' + ('/2' if half else '')
            sheet[f'L{row}'] = f'=IF(C{row}="B",F{row},IF(C{row}="Y",IFERROR(G{row}/M1*11,0),IFERROR(G{row}/M1*12,0)))'
            sheet[f'M{row}'] = f'=L{row}-F{row}'
            sheet[f'N{row}'] = f'=IF(D{row}="X","",M{row})'
    codes = wb.create_sheet('Codes')
    codes.append(['Category 1','Category 2','Category 3','VLookUp Code','BC Code'])
    codes.append(['Synthetic department','Synthetic section','Synthetic line','Synthetic','SYNTHETIC'])
    wb.create_sheet('Actual 2026').append(['Year',2026])
    for index in range(3,sheets): wb.create_sheet(f'Ancillary {index}')
    output = BytesIO(); wb.save(output); wb.close()
    data = output.getvalue()
    if unsized:
        from xml.etree import ElementTree as ET
        def remove_dimension(xml):
            root=ET.fromstring(xml)
            for child in list(root):
                if child.tag.endswith('}dimension'): root.remove(child)
            return ET.tostring(root)
        data = rewrite(data, {f'xl/worksheets/sheet{i}.xml':remove_dimension for i in range(1,sheets+1)})
    # Replay fixtures must be byte-stable across seconds and suite order.
    from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    canonical = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(canonical, 'w') as target:
        for entry in source.infolist():
            value = source.read(entry)
            if entry.filename == 'docProps/core.xml':
                root = ET.fromstring(value)
                for element in root:
                    if element.tag.rsplit('}', 1)[-1] in ('created', 'modified'):
                        element.text = '2026-01-01T00:00:00Z'
                value = ET.tostring(root)
            info = ZipInfo(entry.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, value)
    return canonical.getvalue()


def budget_ledger(user):
    from api.finance_run_test_utils import golden, candidate, approve
    artifact = golden()
    for index, row in enumerate(artifact['ledger']['rows']):
        row['bc'] = 'synthetic' if index % 2 else 'SYNTHETIC'
    return approve(candidate(user, artifact=artifact), user)
