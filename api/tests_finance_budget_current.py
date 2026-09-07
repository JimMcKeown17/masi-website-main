from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from unittest import skipUnless
from unittest.mock import patch
from django.db import connection, connections, transaction
from django.test import TestCase, TransactionTestCase
from rest_framework.test import APIClient
from api.finance_run_test_utils import actor, candidate, approve
from api.finance_budget_test_utils import budget_workbook, budget_ledger
from api.models import FinanceRun
from api.parsers.finance_workbook import MIME
from api.services import finance_runs as service


class BudgetCurrentTests(TestCase):
    def setUp(self):
        self.user=actor(); self.dep=budget_ledger(self.user)
        self.client=APIClient(); self.client.force_authenticate(self.user)

    def budget(self):
        run,status=service.upload_workbook(BytesIO(budget_workbook()),self.user,kind='budgets',year=2026,
                source_name='20260907 - Synthetic.xlsx',content_type=MIME,ledger_run_id=self.dep.pk)
        self.assertEqual(status,201)
        return approve(run,self.user)

    def current(self): return self.client.get('/api/finance/current/?year=2026').data

    def test_same_management_sha_is_compatible_across_kinds(self):
        budget=self.budget(); data=self.current()
        self.assertTrue(data['compatible'])
        self.assertNotEqual(budget.source_sha256,self.dep.source_sha256)
        self.assertEqual(data['runs']['budgets']['management_accounts_sha256'],self.dep.source_sha256)
        self.assertEqual(data['runs']['budgets']['budget_source_sha256'],budget.source_sha256)

    def test_missing_kind_tolerated_but_unresolved_dependency_is_not(self):
        self.assertTrue(self.current()['compatible'])
        budget=self.budget()
        budget.manifest['dependencies']=[]; budget.save(update_fields=['manifest'])
        self.assertFalse(self.current()['compatible'])
        self.assertEqual(self.current()['compatibility_reason']['code'],'DEPENDENCY_UNRESOLVED')

    def test_new_funder_current_marks_old_budget_incompatible_without_recompute(self):
        budget=self.budget(); digest=budget.payload_sha256
        approve(candidate(self.user,sha='b'*64),self.user,override_anti_rollback=True)
        self.assertFalse(self.current()['compatible'])
        budget.refresh_from_db(); self.assertEqual(budget.payload_sha256,digest)
        self.assertEqual(self.client.get(f'/api/finance/runs/{budget.pk}/').status_code,200)

    def test_budget_approval_does_not_change_funder_snapshot_or_years(self):
        before=self.client.get('/api/finance/snapshot/').json()
        self.budget()
        self.assertEqual(before,self.client.get('/api/finance/snapshot/').json())

    def test_contributor_rows_exports_and_authorization_rechecked(self):
        import csv
        from io import StringIO, BytesIO
        from openpyxl import load_workbook
        from django.contrib.auth.models import Permission
        budget=self.budget()
        path=f'/api/finance/runs/{budget.pk}/rows/'
        query='?year=2026&bc=SyNtHeTiC'
        rows=self.client.get(path+query)
        self.assertEqual(rows.status_code,200,rows.data)
        expected=list(self.dep.ledger_rows.filter(year=2026).order_by('date','sheet_row','row_key').values_list('row_key',flat=True))
        self.assertEqual([r['row_key'] for r in rows.data['results']],expected)
        self.assertGreater(len(expected),0)
        exported=self.client.get(path+'export/'+query+'&format=csv')
        self.assertEqual(exported.status_code,200)
        self.assertEqual([r['row_key'] for r in csv.DictReader(StringIO(exported.content.decode()))],expected)
        exported=self.client.get(path+'export/'+query+'&format=xlsx')
        self.assertEqual(exported.status_code,200)
        wb=load_workbook(BytesIO(exported.content),read_only=True,data_only=True)
        self.assertEqual([r[0] for r in list(wb.active.values)[1:]],expected);wb.close()
        reader=actor('read-only',role='STAFF')
        grant=Permission.objects.get(content_type__app_label='api',content_type__model='financerun',codename='read_finance')
        reader.user_permissions.add(grant);self.client.force_authenticate(reader)
        self.assertEqual(self.client.get(path+query).status_code,200)
        reader.user_permissions.clear()
        from django.contrib.auth.models import User
        self.client.force_authenticate(User.objects.get(pk=reader.pk))
        self.assertEqual(self.client.get(path+query).status_code,403)
        self.assertEqual(self.client.get(path+'export/'+query+'&format=csv').status_code,403)
        self.client.force_authenticate(self.user)
        self.dep.manifest['source']['sha256']='f'*64
        self.dep.save(update_fields=['manifest'])
        self.assertEqual(self.client.get(path+query).status_code,400)

    def test_budget_only_current_and_same_sha_distinct_ledger_ids(self):
        self.dep.producer_version='0.2.1'
        self.dep.manifest['producer']['version']='0.2.1'
        self.dep.save(update_fields=['producer_version','manifest'])
        with patch.dict(service.SUPPORTED_PAIRS,{('2.0.0','0.2.1'):'2.0.0'}):
            budget=self.budget()
            self.dep.status='superseded';self.dep.save(update_fields=['status'])
            self.assertTrue(self.current()['compatible'])
            newer=approve(candidate(self.user),self.user)
            self.assertNotEqual(newer.pk,budget.dependency_run_id)
            self.assertEqual(newer.source_sha256,self.dep.source_sha256)
            self.assertTrue(self.current()['compatible'])


@skipUnless(connection.vendor=='postgresql','Requires PostgreSQL advisory locks, row locks and separate connections; SQLite is functional evidence only.')
class BudgetPostgresTests(TransactionTestCase):
    # Migration-created rows (the Finance Managers group) must survive this class's flush;
    # without this, later tests depend on alphabetical ordering.
    serialized_rollback = True
    def setUp(self):
        self.user=actor(); self.dep=budget_ledger(self.user)
        self.data=budget_workbook()

    def upload(self):
        return service.upload_workbook(BytesIO(self.data),self.user,kind='budgets',year=2026,
            source_name='20260907 - Synthetic.xlsx',content_type=MIME,ledger_run_id=self.dep.pk)

    def worker(self,fn):
        connections.close_all()
        try: return fn()
        finally: connections.close_all()

    def test_postgres_racing_approval_upload_and_injected_failure_preserve_current(self):
        from threading import Event
        next_funder=candidate(self.user,sha='b'*64)
        locked,release=Event(),Event()
        def checkpoint(stage):
            if stage=='after_lock':
                locked.set()
                if not release.wait(10): raise RuntimeError('thread timeout')
        with ThreadPoolExecutor(2) as pool, patch.object(service,'_transition_checkpoint',side_effect=checkpoint):
            first=pool.submit(self.worker,lambda:approve(next_funder,self.user,override_anti_rollback=True))
            try:
                self.assertTrue(locked.wait(10))
                with self.assertRaisesRegex(service.FinanceRunError,'UPLOAD_IN_PROGRESS'):
                    pool.submit(self.worker,self.upload).result(10)
            finally: release.set()
            first.result(10)
        a,_=self.upload();approve(a,self.user)
        self.data=budget_workbook(sheets=4)
        b,_=self.upload()
        def fail(stage):
            if stage=='before_promote': raise RuntimeError('injected')
        with patch.object(service,'_transition_checkpoint',side_effect=fail):
            with self.assertRaises(RuntimeError): approve(b,self.user,override_anti_rollback=True)
        a.refresh_from_db();b.refresh_from_db();next_funder.refresh_from_db()
        self.assertEqual((a.status,b.status,next_funder.status),('approved','candidate','approved'))
        approve(b,self.user,override_anti_rollback=True)

    def test_postgres_dependency_admission_uses_consistent_lock_order(self):
        with patch.object(service,'acquire_tuple_lock',wraps=service.acquire_tuple_lock) as lock:
            run,_=self.upload()
            upload_order=[call.args for call in lock.call_args_list]
        with patch.object(service,'acquire_tuple_lock',wraps=service.acquire_tuple_lock) as lock:
            approve(run,self.user)
            approval_order=[call.args for call in lock.call_args_list]
        self.assertEqual(upload_order,[('budgets',2026),('funders',2026)])
        self.assertEqual(approval_order,upload_order)
