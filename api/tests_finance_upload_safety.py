"""Synthetic XLSX and adversarial ZIP tests; no real financial workbooks."""
from datetime import date
from io import BytesIO
import os
import re
import struct
import zipfile
import warnings
from unittest.mock import patch
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NAME = '20260901 - Synthetic.xlsx'


def workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Expenditure'
    ws.append(['Date', 'Year', 'Name', 'Amount', 'Paid By', 'Category 1',
               'Category 2', 'Category 3', 'BC', 'Allocation', 'Budget Key'])
    ws.append([date(2026, 1, 2), 2026, 'Synthetic', 10, 'Bank', 'Programme',
               'Supplies', None, None, 10, 'TEST'])
    ws = wb.create_sheet('Funder Budgets')
    ws.append(['TEST', 'Synthetic funder'])
    ws.append([None, 'Category', 'Budget', 'Spent'])
    ws.append([None, 'Supplies', 100,
               '=SUM(SUMIFS(Expenditure!J:J,Expenditure!G:G,B3,Expenditure!K:K,"TEST"))'])
    ws.append([None, 'Total'])
    ws.append([])
    ws.append(['Budget Key', 'Funder', 'Start Date', 'End Date', 'Description'])
    ws.append(['TEST', 'Synthetic funder', date(2026, 1, 1), date(2026, 12, 31), 'Synthetic'])
    with BytesIO() as output:
        wb.save(output)
        wb.close()
        return output.getvalue()


def rewrite(data, changes=None, extra=()):
    with BytesIO() as output, zipfile.ZipFile(BytesIO(data)) as source:
        with warnings.catch_warnings(), zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
            warnings.filterwarnings('ignore', message='Duplicate name:', category=UserWarning)
            for entry in source.infolist():
                value = source.read(entry)
                if changes and entry.filename in changes:
                    value = changes[entry.filename](value)
                target.writestr(entry.filename, value)
            for name, value in extra:
                target.writestr(name, value)
        return output.getvalue()


class BodyForbidden(BytesIO):
    @property
    def body(self):
        raise AssertionError('request.body is forbidden')


class WorkbookSafetyTests(SimpleTestCase):
    def setUp(self):
        from api.parsers import finance_workbook
        self.p = finance_workbook
        self.data = workbook_bytes()

    def preflight(self, data=None, **kwargs):
        return self.p.preflight(BodyForbidden(data if data is not None else self.data),
                                source_name=NAME, content_type=MIME, **kwargs)

    def reject(self, data, code, *, sheet=False, **kwargs):
        with self.assertRaises(self.p.WorkbookError) as caught:
            with self.preflight(data, **kwargs) as upload:
                if sheet:
                    self.p.scan_workbook(upload)
        self.assertEqual(caught.exception.code, code)
        self.assertEqual(str(caught.exception), code)

    def test_valid_body_never_accesses_body_and_buffer_closes(self):
        with self.preflight() as upload:
            self.p.scan_workbook(upload)
            buffer = upload.buffer
        self.assertTrue(buffer.closed)

    def _assert_literal_header_cells_scan(self, transform=lambda xml: xml):
        def change(xml):
            # Literal Excel-style empty styled cells, after the last populated
            # header and immediately before row 1 closes; row 2 follows.
            xml = xml.replace(b'</row>', b'<c r="BX1" s="2"/>'
                              b'<c r="BY1" s="2"/><c r="BZ1" s="2"/></row>', 1)
            return transform(xml)

        with self.preflight() as upload:
            self.p.scan_workbook(upload)
        data = rewrite(self.data, {'xl/worksheets/sheet1.xml': change})
        with self.preflight(data) as upload:
            self.p.scan_workbook(upload)
        # Empty styled cells must not increase H: a populated L3 still rejects.
        invalid = rewrite(data, {'xl/worksheets/sheet1.xml': lambda xml: xml.replace(
            b'</sheetData>', transform(b'<row r="3"><c r="L3"><v>1</v></c></row>')
            + b'</sheetData>')})
        self.reject(invalid, 'LEDGER_DATA_BEYOND_HEADER', sheet=True)

    def test_trailing_empty_styled_header_cells_scan(self):
        self._assert_literal_header_cells_scan()

    def test_reordered_cell_attributes_scan(self):
        self._assert_literal_header_cells_scan(
            lambda xml: xml.replace(b'<c r="A1"', b'<c s="1" r="A1"'))

    def test_namespace_prefixed_cells_scan(self):
        self._assert_literal_header_cells_scan(lambda xml: xml.replace(
            b'<c ', b'<x:c xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        ).replace(b'</c>', b'</x:c>'))

    def test_compressed_cap_missing_lying_and_true_length_stops_at_plus_one(self):
        for length in (None, '1', '1024'):
            with self.subTest(length=length), patch.object(self.p, 'MAX_COMPRESSED', 1024):
                stream = BodyForbidden(b'x' * 4096)
                with self.assertRaises(self.p.WorkbookError) as caught:
                    with self.p.preflight(stream, source_name=NAME, content_type=MIME, content_length=length):
                        pass
                self.assertEqual(caught.exception.code, 'UPLOAD_TOO_LARGE')
                self.assertEqual(stream.tell(), 1025)

    def test_declared_length_invalid_before_read(self):
        for length in ('0', '-1', 'bad', str(32 * 1024 * 1024 + 1)):
            with self.subTest(length=length):
                self.reject(b'', 'CONTENT_LENGTH_INVALID', content_length=length)

    def test_envelope_type_and_basename(self):
        for field, value, code in [('content_type', 'multipart/form-data', 'CONTENT_TYPE_INVALID'),
                                   ('source_name', '../20260901 - Private.xlsx', 'SOURCE_INVALID'),
                                   ('source_name', '20260230 - Private.xlsx', 'SOURCE_INVALID')]:
            options = dict(source_name=NAME, content_type=MIME)
            options[field] = value
            with self.assertRaises(self.p.WorkbookError) as caught:
                with self.p.preflight(BodyForbidden(self.data), **options):
                    pass
            self.assertEqual(caught.exception.code, code)

    def test_zip_entry_count(self):
        self.reject(rewrite(self.data, extra=[(f'extra/{n}', b'x') for n in range(257)]), 'ZIP_ENTRY_LIMIT')

    def test_zip_duplicate_and_paths(self):
        for name, code in [('xl/workbook.xml', 'ZIP_DUPLICATE_ENTRY'), ('../x', 'ZIP_PATH_INVALID'),
                           ('/x', 'ZIP_PATH_INVALID'), ('a\\x', 'ZIP_PATH_INVALID'),
                           ('C:/x', 'ZIP_PATH_INVALID')]:
            with self.subTest(name=name):
                self.reject(rewrite(self.data, extra=[(name, b'x')]), code)

    def test_zip_encrypted(self):
        data = bytearray(self.data)
        for signature, offset in [(b'PK\x03\x04', 6), (b'PK\x01\x02', 8)]:
            index = data.find(signature)
            flags = struct.unpack_from('<H', data, index + offset)[0]
            struct.pack_into('<H', data, index + offset, flags | 1)
        self.reject(bytes(data), 'ZIP_ENCRYPTED')

    def test_zip_expanded_and_ratio_limits(self):
        with patch.object(self.p, 'MAX_ENTRY_EXPANDED', 100):
            self.reject(self.data, 'ZIP_ENTRY_EXPANDED_LIMIT')
        with patch.object(self.p, 'MAX_TOTAL_EXPANDED', 100):
            self.reject(self.data, 'ZIP_TOTAL_EXPANDED_LIMIT')
        self.reject(rewrite(self.data, extra=[('bomb', b'x' * 100000)]), 'ZIP_RATIO_LIMIT')

    def test_missing_duplicate_sheets_and_headers(self):
        for old, new, code in [(b'name="Expenditure"', b'name="Other"', 'REQUIRED_SHEETS'),
                               (b'name="Funder Budgets"', b'name="Expenditure"', 'REQUIRED_SHEETS')]:
            self.reject(rewrite(self.data, {'xl/workbook.xml': lambda x: x.replace(old, new)}), code, sheet=True)
        for path, old, code in [('xl/worksheets/sheet1.xml', b'>Paid By<', 'LEDGER_REQUIRED_HEADER'),
                                ('xl/worksheets/sheet2.xml', b'>Start Date<', 'CONTRACT_KEY_HEADER')]:
            self.reject(rewrite(self.data, {path: lambda x: x.replace(old, b'>Other<')}), code, sheet=True)

    def test_declared_rows_columns_and_product(self):
        for extent, code in [('A1:K50001', 'SHEET_BOUNDS'), ('A1:IW2', 'SHEET_BOUNDS'),
                              ('A1:DX50000', 'LEDGER_CELL_LIMIT')]:
            def change(x):
                x = re.sub(rb'<dimension ref="[^"]+"', f'<dimension ref="{extent}"'.encode(), x)
                if extent == 'A1:DX50000':
                    x = x.replace(b'</row>', b'<c r="DX1" t="inlineStr"><is><t>Extra</t></is></c></row>', 1)
                return x
            self.reject(rewrite(self.data, {'xl/worksheets/sheet1.xml': change}), code, sheet=True)

    def test_spoofed_dimensions_sparse_headers_and_data_beyond_header(self):
        for cell, code in [('A50001', 'SHEET_BOUNDS'), ('L2', 'LEDGER_DATA_BEYOND_HEADER'),
                           ('DY1', 'LEDGER_HEADER_LIMIT'), ('LCV1', 'SHEET_BOUNDS')]:
            def change(x):
                row = re.search(r'\d+', cell).group()
                from xml.etree import ElementTree as ET
                ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                root = ET.fromstring(x)
                sheet = root.find(ns + 'sheetData')
                existing = sheet.find(ns + 'row' + f"[@r='{row}']")
                target = existing if existing is not None else ET.SubElement(sheet, ns + 'row', r=row)
                value = ET.SubElement(target, ns + 'c', r=cell, t='inlineStr')
                ET.SubElement(ET.SubElement(value, ns + 'is'), ns + 't').text = 'Private'
                return ET.tostring(root)
            self.reject(rewrite(self.data, {'xl/worksheets/sheet1.xml': change}), code, sheet=True)

    def test_defused_xml_rejects_entities(self):
        data = rewrite(self.data, {'xl/workbook.xml': lambda x: b'<!DOCTYPE x [<!ENTITY secret "private">]>' + x})
        self.reject(data, 'XML_INVALID', sheet=True)

    def test_valid_body_above_default_django_memory_limit(self):
        data = rewrite(self.data, extra=[('padding.bin', os.urandom(2700000))])
        self.assertGreater(len(data), 2621440)
        with self.preflight(data) as upload:
            self.p.scan_workbook(upload)

    def test_four_million_cell_boundary_and_budget_bounds(self):
        def boundary(x):
            x = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:DX31250"', x)
            return x.replace(b'</row>', b'<c r="DX1" t="inlineStr"><is><t>Extra</t></is></c></row>', 1)
        data = rewrite(self.data, {'xl/worksheets/sheet1.xml': boundary})
        with self.preflight(data) as upload:
            self.p.scan_workbook(upload)
        self.reject(rewrite(data, {'xl/worksheets/sheet1.xml': lambda x: x.replace(b'DX31250', b'DX31251')}), 'LEDGER_CELL_LIMIT', sheet=True)
        self.reject(rewrite(self.data, {'xl/worksheets/sheet2.xml': lambda x: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:E5001"', x)}), 'SHEET_BOUNDS', sheet=True)

    def test_understated_zip_size_does_not_hide_expansion(self):
        data = bytearray(rewrite(self.data, extra=[('bomb', b'x' * 100000)]))
        start = 0
        while True:
            start = data.find(b'PK\x01\x02', start)
            if start < 0:
                break
            name_length = struct.unpack_from('<H', data, start + 28)[0]
            if data[start + 46:start + 46 + name_length] == b'bomb':
                struct.pack_into('<I', data, start + 24, 1)
                break
            start += 4
        self.reject(bytes(data), 'ZIP_RATIO_LIMIT')

    def test_buffer_closed_when_sheet_scan_fails(self):
        data = rewrite(self.data, {'xl/workbook.xml': lambda x: b'invalid'})
        with self.assertRaises(self.p.WorkbookError):
            with self.preflight(data) as upload:
                self.p.scan_workbook(upload)
        self.assertTrue(upload.buffer.closed)

    def test_actual_32_mib_cap_reads_only_one_excess_byte(self):
        class GeneratedStream:
            consumed = 0
            def read(self, count):
                self.consumed += count
                return b'x' * count
        stream = GeneratedStream()
        with self.assertRaises(self.p.WorkbookError) as caught:
            with self.p.preflight(stream, source_name=NAME, content_type=MIME):
                pass
        self.assertEqual(caught.exception.code, 'UPLOAD_TOO_LARGE')
        self.assertEqual(stream.consumed, 32 * 1024 * 1024 + 1)

    def test_streamed_product_cannot_hide_behind_small_dimensions(self):
        def change(x):
            x = x.replace(b'</row>', b'<c r="DX1" t="inlineStr"><is><t>Extra</t></is></c></row>', 1)
            return x.replace(b'</sheetData>', b'<row r="31251"><c r="A31251"><v>1</v></c></row></sheetData>')
        self.reject(rewrite(self.data, {'xl/worksheets/sheet1.xml': change}), 'LEDGER_CELL_LIMIT', sheet=True)

    def test_alternate_selected_workbook_rejected_before_producer(self):
        with zipfile.ZipFile(BytesIO(self.data)) as archive:
            alternate = archive.read('xl/worksheets/sheet1.xml').replace(
                b'</row>', b'<c r="LCV1" t="inlineStr"><is><t>Hidden</t></is></c></row>', 1)
            data = rewrite(self.data, {'[Content_Types].xml': lambda x: x.replace(
                b'PartName="/xl/workbook.xml"', b'PartName="/xl/alternate.xml"')}, extra=[
                ('xl/alternate.xml', archive.read('xl/workbook.xml')),
                ('xl/_rels/alternate.xml.rels', archive.read('xl/_rels/workbook.xml.rels').replace(
                    b'/xl/worksheets/sheet1.xml', b'/xl/worksheets/alternate.xml')),
                ('xl/worksheets/alternate.xml', alternate)])
        with patch('api.services.finance_runs.build_run_artifact') as producer, \
             patch('openpyxl.load_workbook') as loader:
            self.reject(data, 'WORKBOOK_METADATA_INVALID', sheet=True)
        producer.assert_not_called()
        loader.assert_not_called()

    def test_ambiguous_workbook_and_shared_string_selection_rejected(self):
        workbook_type = b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'
        for declaration in (
            b'<Override PartName="/xl/other.xml" ContentType="' + workbook_type + b'"/>',
            b'<Default Extension="xml" ContentType="' + workbook_type + b'"/>',
            b'<Override PartName="/xl/otherStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>',
        ):
            with self.subTest(declaration=declaration):
                self.reject(rewrite(self.data, {'[Content_Types].xml': lambda x: x.replace(
                    b'</Types>', declaration + b'</Types>')}), 'WORKBOOK_METADATA_INVALID', sheet=True)

    def test_scanned_sheet_paths_match_openpyxl_selected_paths(self):
        import openpyxl
        with self.preflight() as upload, patch.object(self.p, '_scan_sheet', wraps=self.p._scan_sheet) as scan:
            self.p.scan_workbook(upload)
            wb = openpyxl.load_workbook(upload.buffer, read_only=True)
            try:
                self.assertEqual({call.args[1] for call in scan.call_args_list},
                                 {wb[name]._worksheet_path for name in ('Expenditure', 'Funder Budgets')})
            finally:
                wb.close()

    def test_sheet_doctype_or_entity_rejected_before_xml_parser(self):
        for path in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
            for declaration in (b'<!DOCTYPE worksheet>', b'<!ENTITY secret "private">'):
                with self.subTest(path=path, declaration=declaration):
                    data = rewrite(self.data, {path: lambda xml: declaration + xml})
                    with self.preflight(data) as upload, zipfile.ZipFile(upload.buffer) as archive:
                        with patch.object(self.p, 'iterparse', side_effect=AssertionError('XML parser touched sheet')), \
                             patch.object(self.p, 'fromstring', side_effect=AssertionError('XML parser touched sheet')):
                            with self.assertRaises(self.p.WorkbookError) as caught:
                                self.p._scan_sheet(archive, path, 'Expenditure' if 'sheet1' in path else 'Funder Budgets', [])
                        self.assertEqual(caught.exception.code, 'SHEET_XML_DECLARATION')

    def test_sheet_scan_never_reads_whole_member(self):
        path = 'xl/worksheets/sheet1.xml'
        padding = b'<!--' + os.urandom(self.p.CHUNK).hex().encode() + b'-->'
        data = rewrite(self.data, {path: lambda xml: xml.replace(b'</worksheet>', padding + b'</worksheet>')})
        original = zipfile.ZipFile.read
        def bounded_read(archive, name, *args, **kwargs):
            member = name.filename if isinstance(name, zipfile.ZipInfo) else name
            self.assertFalse(member.startswith('xl/worksheets/'), 'whole worksheet read forbidden')
            return original(archive, name, *args, **kwargs)
        with self.preflight(data) as upload, patch.object(zipfile.ZipFile, 'read', bounded_read):
            self.p.scan_workbook(upload)

    def test_split_encoded_declarations_rejected_before_parser(self):
        path = 'xl/worksheets/sheet1.xml'
        for encoding in ('utf-8', 'utf-16-le', 'utf-16-be', 'utf-32-le', 'utf-32-be'):
            for token in ('<!DOCTYPE worksheet>', '<!ENTITY secret "private">'):
                with self.subTest(encoding=encoding, token=token):
                    declaration = token.encode(encoding)
                    # Split inside the declaration token, including its NUL bytes.
                    xml = b' ' * (self.p.CHUNK - 5) + declaration
                    with BytesIO() as buffer:
                        with zipfile.ZipFile(buffer, 'w') as archive:
                            archive.writestr(path, xml)
                        with zipfile.ZipFile(buffer) as archive, patch.object(
                                self.p, 'iterparse', side_effect=AssertionError('parser invoked')):
                            with self.assertRaises(self.p.WorkbookError) as caught:
                                self.p._scan_sheet(archive, path, 'Expenditure', [])
                        self.assertEqual(caught.exception.code, 'SHEET_XML_DECLARATION')

    def test_metadata_part_caps_before_parsing(self):
        for path in ('[Content_Types].xml', 'xl/workbook.xml',
                     'xl/_rels/workbook.xml.rels', '_rels/.rels'):
            with self.subTest(path=path):
                data = rewrite(self.data, {path: lambda xml:
                    xml + b'<!--' + os.urandom(524289).hex().encode() + b'-->'})
                with patch.object(self.p, 'fromstring') as parser:
                    self.reject(data, 'PART_SIZE_LIMIT', sheet=True)
                parser.assert_not_called()

    def test_shared_strings_actual_expansion_cap_with_understated_directory(self):
        path = 'xl/sharedStrings.xml'
        # Generate compressed fixture incrementally; never allocate a 64 MiB part.
        with BytesIO() as buffer:
            with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
                with archive.open(path, 'w') as part:
                    for _ in range(1024):
                        part.write(b' ' * 65536)
                    part.write(b' ')
            data = bytearray(buffer.getvalue())
        with patch.object(zipfile.ZipFile, 'read', side_effect=AssertionError('whole part read')):
            self.reject(bytes(data), 'PART_SIZE_LIMIT')
        directory = data.index(b'PK\x01\x02')
        struct.pack_into('<I', data, directory + 24, 1)
        with patch.object(self.p, 'MAX_RATIO', 100000), patch.object(
                zipfile.ZipFile, 'read', side_effect=AssertionError('whole part read')):
            self.reject(bytes(data), 'PART_SIZE_LIMIT')

    def test_metadata_actual_expansion_cap_with_understated_directory(self):
        data = bytearray(rewrite(self.data, {'[Content_Types].xml': lambda xml:
            xml + b'<!--' + b'x' * (1024 * 1024) + b'-->'}))
        offset = 0
        while True:
            offset = data.index(b'PK\x01\x02', offset)
            if data[offset + 46:offset + 46 + len(b'[Content_Types].xml')] == b'[Content_Types].xml':
                break
            offset += 4
        struct.pack_into('<I', data, offset + 24, 1)
        with patch.object(self.p, 'MAX_RATIO', 100000):
            self.reject(bytes(data), 'PART_SIZE_LIMIT')

    def _strings_archive(self, entries):
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
            with archive.open('xl/sharedStrings.xml', 'w') as part:
                part.write(('<sst xmlns="' + self.p.NS[1:-1] + '">').encode())
                for entry in entries:
                    part.write(entry)
                part.write(b'</sst>')
        return buffer

    def test_shared_string_four_million_count_cap(self):
        from itertools import chain, repeat
        with self._strings_archive(chain(repeat(b'<si/>' * 1000, 4000), [b'<si/>'])) as buffer:
            with zipfile.ZipFile(buffer) as archive, self.assertRaises(self.p.WorkbookError) as caught:
                self.p._shared_strings(archive)
            self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')

    def test_shared_string_length_cap_including_rich_text(self):
        for entry in (b'<si><t>' + b'x' * 32768 + b'</t></si>',
                      b'<si><r><t>' + b'x' * 16384 + b'</t></r><r><t>'
                      + b'x' * 16384 + b'</t></r></si>'):
            with self._strings_archive([entry]) as buffer:
                with zipfile.ZipFile(buffer) as archive, self.assertRaises(self.p.WorkbookError) as caught:
                    self.p._shared_strings(archive)
                self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')

    def test_shared_strings_declaration_before_parser(self):
        with self._strings_archive([b'<!DOCTYPE sst>']) as buffer:
            with zipfile.ZipFile(buffer) as archive, patch.object(self.p, 'iterparse') as parser:
                with self.assertRaises(self.p.WorkbookError) as caught:
                    self.p._shared_strings(archive)
                self.assertEqual(caught.exception.code, 'XML_INVALID')
                parser.assert_not_called()

    def test_normal_shared_strings_stream_and_resolve_headers(self):
        labels = list(self.p.LEDGER_HEADERS)
        strings = ('<sst xmlns="' + self.p.NS[1:-1] + '">' + ''.join(
            '<si><t>' + label + '</t></si>' for label in labels)
            + '<si><t> </t></si><si><r><t>Other</t></r></si></sst>').encode()
        def shared_headers(xml):
            for index, label in enumerate(labels):
                xml = xml.replace(b't="inlineStr"><is><t>' + label.encode() + b'</t></is>',
                                  f't="s"><v>{index}</v>'.encode())
            return xml
        data = rewrite(self.data, {
            '[Content_Types].xml': lambda xml: xml.replace(b'</Types>',
                b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'),
            'xl/worksheets/sheet1.xml': shared_headers,
        }, extra=[('xl/sharedStrings.xml', strings)])
        original = zipfile.ZipFile.read
        def bounded_read(archive, name, *args, **kwargs):
            self.assertNotEqual(getattr(name, 'filename', name), 'xl/sharedStrings.xml')
            return original(archive, name, *args, **kwargs)
        with self.preflight(data) as upload, patch.object(zipfile.ZipFile, 'read', bounded_read):
            self.p.scan_workbook(upload)
            with zipfile.ZipFile(upload.buffer) as archive:
                self.assertEqual(self.p._shared_strings(archive), labels + [False, True])

    def test_relationship_resolution_matches_openpyxl(self):
        import openpyxl
        for target in ('worksheets/sheet1.xml', '/xl/worksheets/sheet1.xml', '/xl/worksheets/Sheet1.xml'):
            with self.subTest(target=target):
                data = rewrite(self.data, {'xl/_rels/workbook.xml.rels': lambda xml: xml.replace(
                    b'/xl/worksheets/sheet1.xml', target.encode())})
                if 'Sheet1' in target:
                    with BytesIO() as output, zipfile.ZipFile(BytesIO(data)) as source:
                        with zipfile.ZipFile(output, 'w') as dest:
                            for entry in source.infolist():
                                dest.writestr(entry.filename.replace('sheet1.xml', 'Sheet1.xml'), source.read(entry))
                        data = output.getvalue()
                with self.preflight(data) as upload, patch.object(self.p, '_scan_sheet', wraps=self.p._scan_sheet) as scan:
                    self.p.scan_workbook(upload)
                    wb = openpyxl.load_workbook(upload.buffer, read_only=True)
                    try:
                        self.assertEqual(scan.call_args_list[0].args[1], wb['Expenditure']._worksheet_path)
                    finally:
                        wb.close()

    def test_noncanonical_relationship_targets_rejected(self):
        for target in ('worksheets/sheet1.xml/', '/xl/worksheets/sheet1.xml/',
                       './worksheets/sheet1.xml', 'worksheets/../worksheets/sheet1.xml',
                       'worksheets//sheet1.xml', 'worksheets\\sheet1.xml',
                       '/xl/worksheets/Sheet1.xml'):
            with self.subTest(target=target):
                data = rewrite(self.data, {'xl/_rels/workbook.xml.rels': lambda xml: xml.replace(
                    b'/xl/worksheets/sheet1.xml', target.encode())})
                self.reject(data, 'WORKBOOK_METADATA_INVALID', sheet=True)

    def test_distinct_zip_names_with_same_normalized_path_rejected(self):
        self.reject(rewrite(self.data, extra=[('other/item', b'x'), ('other/item/', b'x')]),
                    'WORKBOOK_METADATA_INVALID')

    def test_utf16_sheet_declaration_rejected_before_parser(self):
        data = rewrite(self.data, {'xl/worksheets/sheet1.xml': lambda xml:
            ('<!DOCTYPE worksheet>' + xml.decode()).encode('utf-16')})
        self.reject(data, 'SHEET_XML_DECLARATION', sheet=True)


class WorkbookSelectionHTTPTests(TestCase):
    def test_part_and_string_limits_http_before_producer(self):
        from rest_framework.test import APIClient
        from urllib.parse import urlencode
        from api.finance_run_test_utils import actor
        from api.models import FinanceRun
        client = APIClient()
        client.force_authenticate(actor())
        data = workbook_bytes()
        oversized_metadata = rewrite(data, {'[Content_Types].xml': lambda xml:
            xml + b'<!--' + b'x' * (1024 * 1024) + b'-->'})
        strings = (b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   b'<si><t>' + b'x' * 32768 + b'</t></si></sst>')
        oversized_string = rewrite(data, {'[Content_Types].xml': lambda xml: xml.replace(
            b'</Types>', b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')},
            extra=[('xl/sharedStrings.xml', strings)])
        for payload, code in ((oversized_metadata, 'PART_SIZE_LIMIT'),
                              (oversized_string, 'SHARED_STRING_LIMIT')):
            with self.subTest(code=code), patch('api.parsers.finance_workbook.MAX_RATIO', 100000), patch(
                    'api.services.finance_runs.build_run_artifact') as producer, patch(
                    'openpyxl.load_workbook') as loader:
                response = client.post('/api/finance/runs/?' + urlencode(
                    dict(kind='funders', year=2026, source_name=NAME)), payload, content_type=MIME)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data, {'code': code})
                self.assertFalse(FinanceRun.objects.exists())
                producer.assert_not_called()
                loader.assert_not_called()

    def test_trailing_slash_target_http_rejected_before_producer_or_openpyxl(self):
        from rest_framework.test import APIClient
        from urllib.parse import urlencode
        from api.finance_run_test_utils import actor
        from api.models import FinanceRun
        data = workbook_bytes()
        with zipfile.ZipFile(BytesIO(data)) as archive:
            benign = archive.read('xl/worksheets/sheet1.xml')
        data = rewrite(data, {
            'xl/_rels/workbook.xml.rels': lambda xml: xml.replace(
                b'/xl/worksheets/sheet1.xml', b'worksheets/sheet1.xml/'),
            'xl/worksheets/sheet1.xml': lambda xml: xml.replace(
                b'</row>', b'<c r="LCV1" t="inlineStr"><is><t>Hidden</t></is></c></row>', 1),
        }, extra=[('xl/worksheets/sheet1.xml/', benign)])
        client = APIClient()
        client.force_authenticate(actor())
        with patch('api.services.finance_runs.build_run_artifact') as producer, \
             patch('openpyxl.load_workbook') as loader:
            response = client.post('/api/finance/runs/?' + urlencode(
                dict(kind='funders', year=2026, source_name=NAME)), data, content_type=MIME)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data, {'code': 'WORKBOOK_METADATA_INVALID'})
        self.assertFalse(FinanceRun.objects.exists())
        producer.assert_not_called()
        loader.assert_not_called()

    def test_alternate_workbook_http_never_calls_producer_or_openpyxl(self):
        from rest_framework.test import APIClient
        from urllib.parse import urlencode
        from api.finance_run_test_utils import actor
        from api.models import FinanceRun
        data = workbook_bytes()
        with zipfile.ZipFile(BytesIO(data)) as archive:
            data = rewrite(data, {'[Content_Types].xml': lambda x: x.replace(
                b'PartName="/xl/workbook.xml"', b'PartName="/xl/alternate.xml"')}, extra=[
                ('xl/alternate.xml', archive.read('xl/workbook.xml')),
                ('xl/_rels/alternate.xml.rels', archive.read('xl/_rels/workbook.xml.rels').replace(
                    b'/xl/worksheets/sheet1.xml', b'/xl/worksheets/alternate.xml')),
                ('xl/worksheets/alternate.xml', archive.read('xl/worksheets/sheet1.xml').replace(
                    b'</row>', b'<c r="LCV1" t="inlineStr"><is><t>Hidden</t></is></c></row>', 1))])
        client = APIClient()
        client.force_authenticate(actor())
        with patch('api.services.finance_runs.build_run_artifact') as producer, \
             patch('openpyxl.load_workbook') as loader:
            response = client.post('/api/finance/runs/?' + urlencode(
                dict(kind='funders', year=2026, source_name=NAME)), data, content_type=MIME)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data, {'code': 'WORKBOOK_METADATA_INVALID'})
        producer.assert_not_called()
        loader.assert_not_called()
        self.assertFalse(FinanceRun.objects.exists())


def string_complexity_workbook(chunks):
    """Hand-written ZIP_STORED XLSX, streamed into memory without disk fixtures."""
    from api.parsers import finance_workbook as p
    ns = p.NS[1:-1]
    rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with BytesIO() as buffer:
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_STORED) as archive:
            archive.writestr('[Content_Types].xml',
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            archive.writestr('xl/workbook.xml', f'<workbook xmlns="{ns}" xmlns:r="{rel}"><sheets>'
                '<sheet name="Expenditure" sheetId="1" r:id="s1"/>'
                '<sheet name="Funder Budgets" sheetId="2" r:id="s2"/></sheets></workbook>')
            archive.writestr('xl/_rels/workbook.xml.rels',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + ''.join(f'<Relationship Id="s{i}" Type="{rel}/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in (1, 2))
                + '</Relationships>')
            for i, headers in enumerate((p.LEDGER_HEADERS, p.CONTRACT_HEADERS), 1):
                archive.writestr(f'xl/worksheets/sheet{i}.xml', f'<worksheet xmlns="{ns}"><sheetData><row r="1">'
                    + ''.join(f'<c r="{chr(65+j)}1" t="inlineStr"><is><t>{h}</t></is></c>' for j, h in enumerate(headers))
                    + '</row></sheetData></worksheet>')
            with archive.open('xl/sharedStrings.xml', 'w') as part:
                part.write(f'<sst xmlns="{ns}">'.encode())
                for chunk in chunks:
                    part.write(chunk)
                part.write(b'</sst>')
        return buffer.getvalue()


def rejected_string_rss_probe():
    import resource
    import sys
    from itertools import chain, repeat
    from api.parsers import finance_workbook as p
    data = string_complexity_workbook(chain([b'<si>'], repeat(b'<r><t/></r>' * 1000, 2200), [b'</si>']))
    original = p.iterparse
    consumed = 0
    def counted(*args, **kwargs):
        nonlocal consumed
        for item in original(*args, **kwargs):
            consumed += 1
            # RED fails safely instead of materializing millions of elements.
            assert consumed <= 2060, 'shared-string rejection consumed too many events'
            yield item
    with patch.object(p, 'iterparse', counted):
        try:
            with p.preflight(BytesIO(data), source_name=NAME, content_type=MIME) as upload:
                p.scan_workbook(upload)
        except p.WorkbookError as error:
            assert error.code == 'SHARED_STRING_LIMIT', error.code
        else:
            raise AssertionError('complex string accepted')
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    peak = peak if sys.platform == 'darwin' else peak * 1024
    print(f'REJECTED_STRING_PEAK_RSS_BYTES={peak}; EVENTS={consumed}', flush=True)
    assert peak < 200 * 1024 * 1024, peak


class SharedStringComplexityTests(SimpleTestCase):
    def scan(self, entries):
        from api.parsers import finance_workbook as p
        data = string_complexity_workbook(entries)
        with p.preflight(BytesIO(data), source_name=NAME, content_type=MIME) as upload:
            p.scan_workbook(upload)

    def test_millions_empty_runs_rejected_early_under_200_mib(self):
        import subprocess
        import sys
        result = subprocess.run([sys.executable, '-c',
            'from api.tests_finance_upload_safety import rejected_string_rss_probe; rejected_string_rss_probe()'],
            capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        print(result.stdout.strip())

    def test_1024_empty_runs_accepted_1025_rejected(self):
        from api.parsers import finance_workbook as p
        self.scan([b'<si>' + b'<r/>' * 1024 + b'</si>'])
        with self.assertRaises(p.WorkbookError) as caught:
            self.scan([b'<si>' + b'<r/>' * 1025 + b'</si>'])
        self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')

    def test_all_descendants_count_together(self):
        from api.parsers import finance_workbook as p
        self.scan([b'<si>' + b'<r><t/></r>' * 512 + b'</si>'])
        for extra in (b'<t/>', b'<rPh/>', b'<phoneticPr/>', b'<rPr><b/></rPr>'):
            with self.subTest(extra=extra), self.assertRaises(p.WorkbookError) as caught:
                self.scan([b'<si>' + b'<r><t/></r>' * 512 + extra + b'</si>'])
            self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')

    def test_text_length_rejected_at_offending_t_end(self):
        from api.parsers import finance_workbook as p
        original = p.iterparse
        last = None
        def counted(*args, **kwargs):
            nonlocal last
            for event, element in original(*args, **kwargs):
                last = (event, element.tag)
                yield event, element
        with patch.object(p, 'iterparse', counted), self.assertRaises(p.WorkbookError) as caught:
            self.scan([b'<si><r><t>' + b'x' * 16384 + b'</t></r><r><t>' + b'x' * 16384 + b'</t></r></si>'])
        self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')
        self.assertEqual(last, ('end', p.NS + 't'))

    def test_part_node_budget_includes_root_and_every_element(self):
        from api.parsers import finance_workbook as p
        self.assertEqual(getattr(p, 'MAX_SHARED_STRING_NODES', None), 8000000)
        with patch.object(p, 'MAX_SHARED_STRING_NODES', 5):
            self.scan([b'<si><t>A</t></si><si><t>B</t></si>'])
            with self.assertRaises(p.WorkbookError) as caught:
                self.scan([b'<si><t>A</t></si><si><t>B</t></si><si/>'])
            self.assertEqual(caught.exception.code, 'SHARED_STRING_LIMIT')

    def test_completed_elements_cleared_and_detached(self):
        from api.parsers import finance_workbook as p
        stream = BytesIO(f'<sst xmlns="{p.NS[1:-1]}"><si><r><t>abc</t></r></si></sst>'.encode())
        stack = []
        ended = None
        for event, element in p._xml_events(stream):
            if ended is not None:
                previous, parent = ended
                self.assertEqual(len(previous), 0)
                self.assertIsNone(previous.text)
                if parent is not None:
                    self.assertNotIn(previous, parent)
            ended = None
            if event == 'start':
                stack.append(element)
            else:
                stack.pop()
                ended = element, stack[-1] if stack else None

    def test_plain_strings_and_accepted_string_heavy_openpyxl_under_512_mib(self):
        import subprocess
        import sys
        result = subprocess.run([sys.executable, '-c', '''
import resource, sys
from itertools import chain, repeat
from io import BytesIO
from api.tests_finance_upload_safety import string_complexity_workbook, NAME, MIME
from api.parsers import finance_workbook as p
from openpyxl import load_workbook
# 24 MB of plain strings plus entries at the accepted rich-text node boundary.
data = string_complexity_workbook(chain(repeat(b'<si><t>' + b'x' * 24000 + b'</t></si>', 1000),
    repeat(b'<si>' + b'<r><t>x</t></r>' * 512 + b'</si>', 100)))
with p.preflight(BytesIO(data), source_name=NAME, content_type=MIME) as upload:
    p.scan_workbook(upload)
    wb = load_workbook(upload.buffer, read_only=True)
    assert len(wb['Expenditure']._shared_strings) == 1100
    assert wb['Expenditure']._shared_strings[0] == 'x' * 24000
    assert wb['Expenditure']._shared_strings[-1] == 'x' * 512
    wb.close()
peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
peak = peak if sys.platform == 'darwin' else peak * 1024
print(f'ACCEPTED_STRING_PEAK_RSS_BYTES={peak}', flush=True)
assert peak < 512 * 1024 * 1024, peak
'''], capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        print(result.stdout.strip())


class SharedStringComplexityHTTPTests(TestCase):
    def test_node_limit_http_400_no_run_producer_or_openpyxl(self):
        from rest_framework.test import APIClient
        from urllib.parse import urlencode
        from api.finance_run_test_utils import actor
        from api.models import FinanceRun
        client = APIClient()
        client.force_authenticate(actor())
        data = string_complexity_workbook([b'<si>' + b'<r><t/></r>' * 513 + b'</si>'])
        with patch('api.services.finance_runs.build_run_artifact') as producer, patch('openpyxl.load_workbook') as loader:
            response = client.post('/api/finance/runs/?' + urlencode(
                dict(kind='funders', year=2026, source_name=NAME)), data, content_type=MIME)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data, {'code': 'SHARED_STRING_LIMIT'})
        self.assertFalse(FinanceRun.objects.exists())
        producer.assert_not_called()
        loader.assert_not_called()


def rejected_out_of_entry_rss_probe():
    import resource
    import sys
    from itertools import repeat
    from api.parsers import finance_workbook as p
    data = string_complexity_workbook(repeat(b'<r><t/></r>' * 1000, 2200))
    assert len(data) == 24202796, len(data)
    original = p.iterparse
    consumed = 0

    def counted(*args, **kwargs):
        nonlocal consumed
        for item in original(*args, **kwargs):
            consumed += 1
            assert consumed <= 2, 'out-of-entry rejection descended into stray subtree'
            yield item

    with patch.object(p, 'iterparse', counted), patch('openpyxl.load_workbook') as loader:
        try:
            with p.preflight(BytesIO(data), source_name=NAME, content_type=MIME) as upload:
                p.scan_workbook(upload)
        except p.WorkbookError as error:
            assert str(error) == error.code == 'SHARED_STRING_LIMIT'
        else:
            raise AssertionError('out-of-entry strings accepted')
        loader.assert_not_called()
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    peak = peak if sys.platform == 'darwin' else peak * 1024
    print(f'OUT_OF_ENTRY_PEAK_RSS_BYTES={peak}; EVENTS={consumed}; UPLOAD_BYTES={len(data)}', flush=True)
    assert peak < 200 * 1024 * 1024, peak


class Round6StructureTests(SimpleTestCase):
    def scan(self, data):
        from api.parsers import finance_workbook as p
        with p.preflight(BytesIO(data), source_name=NAME, content_type=MIME) as upload:
            p.scan_workbook(upload)

    def assert_early_rejection(self, data, code, events):
        from api.parsers import finance_workbook as p
        original = p.iterparse
        consumed = 0

        def counted(*args, **kwargs):
            nonlocal consumed
            for item in original(*args, **kwargs):
                consumed += 1
                self.assertLessEqual(consumed, events)
                yield item

        with patch.object(p, 'iterparse', counted), self.assertRaises(p.WorkbookError) as caught:
            self.scan(data)
        self.assertEqual(str(caught.exception), code)
        self.assertEqual(consumed, events)

    def test_exact_out_of_entry_fixture_rejected_at_second_event_under_200_mib(self):
        import subprocess
        import sys
        result = subprocess.run([sys.executable, '-c',
            'from api.tests_finance_upload_safety import rejected_out_of_entry_rss_probe; rejected_out_of_entry_rss_probe()'],
            capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        print(result.stdout.strip())

    def test_shared_string_root_must_be_canonical_sst(self):
        data = string_complexity_workbook([b'<si><t>Date</t></si>'])
        for transform in (
                lambda x: x.replace(b'sst', b'wrapper'),
                lambda x: x.replace(b'spreadsheetml/2006/main', b'other'),
                lambda x: x.replace(b'<sst xmlns=', b'<si xmlns=').replace(b'</sst>', b'</si>')):
            with self.subTest(transform=transform):
                self.assert_early_rejection(rewrite(data, {'xl/sharedStrings.xml': transform}),
                                            'SHARED_STRING_LIMIT', 1)

    def test_nested_wrapper_outside_entries_rejected_before_descending(self):
        self.assert_early_rejection(string_complexity_workbook([
            b'<wrapper><si><t>private-value</t></si></wrapper>']), 'SHARED_STRING_LIMIT', 2)

    def test_stray_element_between_valid_entries_rejected_on_start(self):
        self.assert_early_rejection(string_complexity_workbook([
            b'<si><t>Date</t></si><stray secret="private-value"/><si><t>Year</t></si>']),
            'SHARED_STRING_LIMIT', 6)

    def test_plain_shared_string_part_and_real_openpyxl_path_unchanged(self):
        import openpyxl
        from api.parsers import finance_workbook as p
        data = string_complexity_workbook([b'<si><t>Date</t></si><si><t>private-value</t></si>'])
        self.scan(data)
        with zipfile.ZipFile(BytesIO(data)) as archive:
            self.assertEqual(p._shared_strings(archive), ['Date', True])
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
        try:
            self.assertEqual(wb['Expenditure']._shared_strings, ['Date', 'private-value'])
            self.assertEqual(next(wb['Expenditure'].values), p.LEDGER_HEADERS)
        finally:
            wb.close()
        self.scan(workbook_bytes())

    def test_worksheet_root_must_be_canonical_before_descending(self):
        data = workbook_bytes()
        for transform in (lambda x: x.replace(b'worksheet', b'wrapper'),
                          lambda x: x.replace(b'spreadsheetml/2006/main', b'other')):
            with self.subTest(transform=transform):
                self.assert_early_rejection(rewrite(data, {'xl/worksheets/sheet1.xml': transform}),
                                            'XML_INVALID', 1)

    def test_worksheet_part_budget_counts_unmodeled_elements_on_start(self):
        from api.parsers import finance_workbook as p
        self.assertEqual(getattr(p, 'MAX_SHEET_NODES', None), 8000000)
        data = workbook_bytes()
        with zipfile.ZipFile(BytesIO(data)) as archive:
            xml = archive.read('xl/worksheets/sheet1.xml')
        nodes = sum(event == 'start' for event, _ in p._xml_events(BytesIO(xml)))
        with patch.object(p, 'MAX_SHEET_NODES', nodes):
            with zipfile.ZipFile(BytesIO(data)) as archive:
                p._scan_sheet(archive, 'xl/worksheets/sheet1.xml', 'Expenditure', [])
            changed = rewrite(data, {'xl/worksheets/sheet1.xml': lambda x:
                x.replace(b'</worksheet>', b'<unmodeled/></worksheet>')})
            with zipfile.ZipFile(BytesIO(changed)) as archive, self.assertRaises(p.WorkbookError) as caught:
                p._scan_sheet(archive, 'xl/worksheets/sheet1.xml', 'Expenditure', [])
            self.assertEqual(str(caught.exception), 'XML_INVALID')

    def test_worksheet_retained_metadata_and_row_subtrees_bounded_on_start(self):
        from api.parsers import finance_workbook as p
        self.assertEqual(getattr(p, 'MAX_SHEET_RETAINED_NODES', None), 65536)
        # Unknown nested wrappers and siblings survive openpyxl's row clearing;
        # recognized metadata is also built before its end-event dispatch.
        for prefix in (b'<wrapper><a><b><c/></b></a></wrapper>',
                       b'<a/><b/><d/><e/>', b'<mergeCells><a><b><d/></b></a></mergeCells>',
                       b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t/></is></c></row></sheetData>'):
            data = rewrite(workbook_bytes(), {'xl/worksheets/sheet1.xml': lambda x:
                x.replace(b'<sheetPr>', prefix + b'<sheetPr>', 1)})
            with self.subTest(prefix=prefix), patch.object(p, 'MAX_SHEET_RETAINED_NODES', 3):
                # Root counts toward cumulative metadata; rows get their own cap.
                self.assert_early_rejection(data, 'XML_INVALID',
                    6 if prefix.startswith(b'<sheetData>') else 6 if prefix.startswith(b'<a/>') else 4)

    def test_worksheet_retained_budget_includes_cleared_row_shells(self):
        from api.parsers import finance_workbook as p
        data = string_complexity_workbook([])
        # First sheet: worksheet + sheetData + one row = three retained nodes.
        with patch.object(p, 'MAX_SHEET_RETAINED_NODES', 40):
            self.scan(data)  # Includes the separate per-row descendant budget.
            data = rewrite(data, {'xl/worksheets/sheet1.xml': lambda x:
                x.replace(b'</sheetData>', b''.join(f'<row r="{row}"/>'.encode() for row in range(2, 40)) + b'</sheetData>')})
            with self.assertRaises(p.WorkbookError) as caught:
                self.scan(data)
            self.assertEqual(str(caught.exception), 'XML_INVALID')


class Round6StructureHTTPTests(TestCase):
    def test_exact_out_of_entry_fixture_never_calls_producer_or_openpyxl(self):
        from itertools import repeat
        from rest_framework.test import APIClient
        from urllib.parse import urlencode
        from api.finance_run_test_utils import actor
        from api.models import FinanceRun
        from api.parsers import finance_workbook as p
        data = string_complexity_workbook(repeat(b'<r><t/></r>' * 1000, 2200))
        self.assertEqual(len(data), 24202796)
        client = APIClient()
        client.force_authenticate(actor())
        original = p.iterparse
        consumed = 0

        def counted(*args, **kwargs):
            nonlocal consumed
            for item in original(*args, **kwargs):
                consumed += 1
                self.assertLessEqual(consumed, 2)
                yield item

        with patch.object(p, 'iterparse', counted), \
             patch('api.services.finance_runs.build_run_artifact') as producer, \
             patch('openpyxl.load_workbook') as loader:
            response = client.post('/api/finance/runs/?' + urlencode(
                dict(kind='funders', year=2026, source_name=NAME)), data, content_type=MIME)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data, {'code': 'SHARED_STRING_LIMIT'})
        self.assertEqual(consumed, 2)
        self.assertFalse(FinanceRun.objects.exists())
        producer.assert_not_called()
        loader.assert_not_called()
