from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from urllib.parse import urlencode
from django.test import TestCase
from rest_framework.test import APIClient
from masi_finance.publish.budget_run import budget_payload_digest
from api.finance_run_test_utils import actor, candidate, approve
from api.finance_budget_test_utils import budget_workbook, budget_ledger
from api.models import FinanceRun
from api.services import finance_runs as service
from api.parsers.finance_workbook import MIME


class BudgetTests(TestCase):
    def setUp(self):
        self.user=actor()
        self.dependency=budget_ledger(self.user)
        self.client=APIClient(); self.client.force_authenticate(self.user)
        self.data=budget_workbook(half=True)

    def upload(self, *, dependency=None, data=None, name='20260907 - Synthetic.xlsx'):
        query=dict(kind='budgets',year=2026,source_name=name,
                   ledger_run_id=str((dependency or self.dependency).pk))
        return self.client.post('/api/finance/runs/?'+urlencode(query),
                                self.data if data is None else data, content_type=MIME)

    def budget(self, **kwargs):
        response=self.upload(**kwargs)
        self.assertEqual(response.status_code,201,response.data)
        run=FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status,'candidate',response.data)
        return run

    def test_budget_raw_upload_reuses_candidate_transaction(self):
        run=self.budget()
        self.assertEqual(run.dependency_run_id,self.dependency.pk)
        self.assertFalse(run.ledger_rows.exists())
        self.assertIsNone(run.facts_sha256)
        self.assertGreater(run.peak_memory_bytes,0)
        with patch.object(service,'validate_stored_run',side_effect=RuntimeError('private')):
            response=self.upload(name='20260908 - Synthetic.xlsx',data=budget_workbook())
        self.assertEqual(response.status_code,500)
        self.assertEqual(FinanceRun.objects.filter(kind='budgets').count(),1)

    def test_worksheet_hyperlinks_preserve_candidate_figures(self):
        from openpyxl import load_workbook
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        from api.parsers import finance_workbook as parser
        control = self.budget()
        workbook = load_workbook(BytesIO(self.data))
        workbook['2026 Budget']['F6'].hyperlink = 'https://example.invalid/doc'
        workbook.create_sheet('Ancillary')['A1'] = 'Supporting document'
        workbook['Ancillary']['A1'].hyperlink = 'https://example.invalid/doc'
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        with ZipFile(BytesIO(data)) as archive:
            for path in ('xl/worksheets/_rels/sheet1.xml.rels',
                         'xl/worksheets/_rels/sheet4.xml.rels'):
                relationships = list(ET.fromstring(archive.read(path)))
                self.assertEqual(len(relationships), 1)
                self.assertEqual(relationships[0].get('TargetMode'), 'External')
                self.assertEqual(relationships[0].get('Type'),
                                 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink')
        with parser.preflight(BytesIO(data), source_name='20260907 - Synthetic.xlsx',
                              content_type=MIME) as upload:
            parser.scan_workbook(upload, kind='budgets', year=2026)
        with patch.object(service, 'build_budget_run_artifact',
                          wraps=service.build_budget_run_artifact) as producer:
            linked = self.budget(data=data)
        producer.assert_called_once()
        self.assertEqual(linked.payload, control.payload)
        self.assertNotIn('https://example.invalid/doc', str(linked.payload))

    def workbook_with_excel_error(self, sheet_name, coordinate):
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(self.data))
        sheet = (workbook[sheet_name] if sheet_name in workbook.sheetnames
                 else workbook.create_sheet(sheet_name))
        sheet[coordinate] = '#NAME?'
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        checked = load_workbook(BytesIO(data), read_only=True)
        try:
            self.assertEqual(checked[sheet_name][coordinate].data_type, 'e')
        finally:
            checked.close()
        return data

    def assert_excel_error_preserves_candidate(self, sheet_name, coordinate):
        control = self.budget()
        data = self.workbook_with_excel_error(sheet_name, coordinate)
        with patch.object(service, 'build_budget_run_artifact',
                          wraps=service.build_budget_run_artifact) as producer:
            response = self.upload(data=data)
        self.assertEqual(response.status_code, 201, response.data)
        producer.assert_called_once()
        run = FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status, 'candidate')
        self.assertIsNone(run.failure)
        self.assertEqual(run.payload, control.payload)

    def test_actual_label_excel_error_preserves_candidate_payload(self):
        self.assert_excel_error_preserves_candidate('Actual 2026', 'A130')

    def test_ancillary_excel_error_preserves_candidate_payload(self):
        self.assert_excel_error_preserves_candidate('Ancillary', 'A1')

    def test_budget_amount_excel_error_records_producer_failure(self):
        data = self.workbook_with_excel_error('2026 Budget', 'F6')
        with patch.object(service, 'build_budget_run_artifact',
                          wraps=service.build_budget_run_artifact) as producer:
            response = self.upload(data=data)
        self.assertEqual(response.status_code, 201, response.data)
        producer.assert_called_once()
        run = FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status, 'failed')
        self.assertEqual(run.failure, {'phase': 'producer',
                                      'code': 'BUDGET_AMOUNT_INVALID',
                                      'message': 'BUDGET_AMOUNT_INVALID'})
        self.assertEqual(run.dependency_run_id, self.dependency.pk)
        self.assertIsNone(run.payload)

    def test_external_reference_refusals_preserve_history_before_producer(self):
        from api.tests_finance_upload_safety import rewrite
        relationship = (b'<Relationship Id="externalTest" Target="https://example.invalid/doc" '
                        b'TargetMode="External" Type="http://schemas.openxmlformats.org/'
                        b'officeDocument/2006/relationships/%s"/>')
        def rel_part(kind):
            return (b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    + relationship % kind + b'</Relationships>')
        cases = {
            'worksheet image': rewrite(self.data, extra=[
                ('xl/worksheets/_rels/sheet1.xml.rels', rel_part(b'image'))]),
            'worksheet unknown type': rewrite(self.data, extra=[
                ('xl/worksheets/_rels/sheet1.xml.rels', rel_part(b'unknown'))]),
            'workbook hyperlink': rewrite(self.data, {'xl/_rels/workbook.xml.rels':
                lambda x: x.replace(b'</Relationships>', relationship % b'hyperlink' + b'</Relationships>')}),
            'nested non-worksheet hyperlink': rewrite(self.data, extra=[
                ('xl/worksheets/_rels/nested/sheet1.xml.rels', rel_part(b'hyperlink'))]),
            'externalLinks part': rewrite(self.data, extra=[('xl/externalLinks/externalLink1.xml', b'<externalLink/>')]),
            'VBA part': rewrite(self.data, extra=[('xl/vbaProject.bin', b'synthetic')]),
            'external formula': rewrite(self.data, {'xl/worksheets/sheet1.xml':
                lambda x: x.replace(b'</sheetData>', b'<row r="10"><c r="F10"><f>\'[Book]Sheet\'!A1</f></c></row></sheetData>')}),
        }
        before = set(FinanceRun.objects.values_list('pk', flat=True))
        for name, data in cases.items():
            with self.subTest(name=name), patch.object(service, 'build_budget_run_artifact') as producer:
                response = self.upload(data=data)
                self.assertEqual(response.status_code, 400, response.data)
                self.assertEqual(response.data['code'], 'BUDGET_EXTERNAL_REFERENCE')
                producer.assert_not_called()
                self.assertEqual(set(FinanceRun.objects.values_list('pk', flat=True)), before)

    def test_same_bytes_new_dependency_is_new_candidate(self):
        a=self.budget()
        dep=approve(candidate(self.user,sha='b'*64),self.user,override_anti_rollback=True)
        b=self.budget(dependency=dep)
        self.assertNotEqual(a.pk,b.pk)
        self.assertEqual(a.source_sha256,b.source_sha256)

    def test_replay_preserves_uploader_and_status(self):
        run=approve(self.budget(),self.user)
        self.client.force_authenticate(actor('second'))
        with patch.object(service,'build_budget_run_artifact',side_effect=AssertionError('replay')):
            response=self.upload()
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['id'],str(run.pk))
        self.assertEqual(response.data['status'],'approved')
        self.assertEqual(response.data['uploaded_by'],self.user.pk)

    def test_dependency_factless_missing_corrupt_candidate_refuses(self):
        import uuid
        pending=candidate(self.user,sha='c'*64)
        for dep in (pending, type('Missing',(),{'pk':uuid.uuid4()})()):
            self.assertEqual(self.upload(dependency=dep).status_code,400)
        self.dependency.ledger_rows.all().delete()
        self.assertEqual(self.upload().status_code,400)
        self.assertFalse(FinanceRun.objects.filter(kind='budgets').exists())

    def test_dependency_fk_and_manifest_must_agree(self):
        run=self.budget()
        run.manifest['dependencies'][0]['run_id']='00000000-0000-0000-0000-000000000001'
        run.payload_sha256=budget_payload_digest(dict(kind='budgets',schema_version=run.schema_version,manifest=run.manifest,derived=run.payload))
        run.save(update_fields=['manifest','payload_sha256'])
        with self.assertRaises(service.FinanceRunError): approve(run,self.user)

    def test_budget_approval_revalidates_derived_and_dependency(self):
        run=self.budget(); original=deepcopy(run.payload)
        mutations=[lambda d:d['lines'][0].update(projected='999.00'),
                   lambda d:d['lines'][0]['calculation_inputs']['budget_assertion'].update(coefficient='01'),
                   lambda d:d['lines'][0]['calculation_inputs']['budget_assertion'].update(scale=-1),
                   lambda d:d['projection'].update(calculation_precision=38),
                   lambda d:d['lines'][0]['monetary_projections']['variance_all']['part_refs'][0].update(sign=-1),
                   lambda d:d['lines'][0]['monetary_projections']['variance_all'].update(residual='1.00')]
        for mutate in mutations:
            run.payload=deepcopy(original); mutate(run.payload)
            run.payload_sha256=budget_payload_digest(dict(kind='budgets',schema_version=run.schema_version,manifest=run.manifest,derived=run.payload))
            run.save(update_fields=['payload','payload_sha256'])
            with self.assertRaises(service.FinanceRunError): approve(run,self.user)
        run.payload=original
        run.payload_sha256=budget_payload_digest(dict(kind='budgets',schema_version=run.schema_version,manifest=run.manifest,derived=original))
        run.save(update_fields=['payload','payload_sha256'])
        with patch('openpyxl.load_workbook',side_effect=AssertionError('retained inputs only')):
            self.assertEqual(approve(run,self.user).status,'approved')
        self.assertEqual(original['lines'][0]['calculation_inputs']['budget_assertion'],{'coefficient':'12345','scale':4})
        self.assertEqual(original['lines'][0]['actual_share'],'1/2')

    def test_funders_retained_candidate_still_approves_after_pin_update(self):
        self.budget()
        with patch.object(service,'UPLOAD_PRODUCER','9.9.9'):
            retained=approve(candidate(self.user,sha='d'*64),self.user,override_anti_rollback=True)
        self.assertEqual(retained.status,'approved')

    def test_demote_replay_reapprove_preserves_acyclic_recovery(self):
        a=approve(self.budget(),self.user)
        b=approve(self.budget(data=budget_workbook(),name='20260908 - Synthetic.xlsx'),self.user)
        restored=service.demote_run(b.pk,self.user,override_anti_rollback=True,acknowledge_findings=True,note='Synthetic recovery')
        self.assertEqual(restored.pk,a.pk)
        replay=self.upload(data=budget_workbook(),name='20260908 - Synthetic.xlsx')
        self.assertEqual(replay.status_code,200)
        self.assertEqual(replay.data['status'],'superseded')
        b=approve(b,self.user)
        a.refresh_from_db()
        self.assertIsNone(a.previous_approved_id)
        self.assertEqual(b.previous_approved_id,a.pk)

    def test_all_mutations_require_publish_and_actors_are_server_derived(self):
        run=self.budget()
        response=self.client.post(f'/api/finance/runs/{run.pk}/approve/',{'approved_by':self.user.pk},format='json')
        self.assertEqual(response.status_code,400)
        reader=actor('staff',role='STAFF'); self.client.force_authenticate(reader)
        self.assertEqual(self.upload().status_code,403)
        for action in ('approve','demote'):
            self.assertEqual(self.client.post(f'/api/finance/runs/{run.pk}/{action}/',{},format='json').status_code,403)
        with self.assertRaises(service.FinanceRunError):
            service.upload_workbook(BytesIO(self.data),reader,kind='budgets',year=2026,source_name='20260907 - Synthetic.xlsx',content_type=MIME,ledger_run_id=self.dependency.pk)

    def test_failed_budget_retains_dependency_and_is_terminal(self):
        from masi_finance.publish.budget_run import BudgetRunError
        with patch.object(service,'build_budget_run_artifact',side_effect=BudgetRunError('BUDGET_HIERARCHY_INVALID')):
            response=self.upload()
        self.assertEqual(response.status_code,201,response.data)
        run=FinanceRun.objects.get(pk=response.data['id'])
        self.assertEqual(run.status,'failed')
        self.assertEqual(run.dependency_run_id,self.dependency.pk)
        self.assertIsNone(run.payload)
        self.assertEqual(self.upload().data['id'],str(run.pk))
        with self.assertRaises(service.FinanceRunError): approve(run,self.user,override_anti_rollback=True)

    def test_budget_constraints_and_protected_dependency(self):
        from django.db import IntegrityError, transaction
        from django.db.models.deletion import ProtectedError
        from django.core.exceptions import ValidationError
        run=self.budget()
        for changes in ({'dependency_run':None},{'facts_sha256':'a'*64},{'fact_row_count':1}):
            with self.assertRaises(IntegrityError), transaction.atomic():
                FinanceRun.objects.filter(pk=run.pk).update(**changes)
        with self.assertRaises(ProtectedError): self.dependency.delete()
        with self.assertRaises(IntegrityError), transaction.atomic():
            FinanceRun.objects.filter(pk=self.dependency.pk).update(dependency_run=run)
        run.dependency_run=run
        with self.assertRaises(ValidationError): run.clean()
        run.dependency_run=self.dependency; run.accounting_year=2025
        with self.assertRaises(ValidationError): run.clean()

    def test_budget_acknowledgement_and_all_three_rollback_guards(self):
        a=self.budget()
        # Use the source reader to produce a real in-scope error, not a manually
        # inserted finding. A missing BC never joins a zero-code group.
        from api.tests_finance_upload_safety import rewrite
        data=rewrite(budget_workbook(),{'xl/worksheets/sheet1.xml':lambda x:x.replace(b'SYNTHETIC',b'')})
        erroneous=self.budget(data=data)
        self.assertGreater(erroneous.in_scope_error_count,0)
        with self.assertRaisesRegex(service.FinanceRunError,'FINDINGS_ACKNOWLEDGEMENT_REQUIRED'):
            service.approve_run(erroneous.pk,self.user)
        with self.assertRaisesRegex(service.FinanceRunError,'NOTE_REQUIRED'):
            service.approve_run(erroneous.pk,self.user,acknowledge_findings=True)
        approve(a,self.user)
        # Same date, different bytes.
        b=self.budget(data=budget_workbook())
        with self.assertRaisesRegex(service.FinanceRunError,'ANTI_ROLLBACK'): approve(b,self.user)
        approve(b,self.user,override_anti_rollback=True)
        # Older source date, distinct bytes to avoid idempotent replay.
        older=self.budget(data=budget_workbook(sheets=4),name='20260906 - Synthetic.xlsx')
        with self.assertRaisesRegex(service.FinanceRunError,'ANTI_ROLLBACK'): approve(older,self.user)
        # Identical bytes, different named ledger, same schema major.
        dep=approve(candidate(self.user,sha='e'*64),self.user,override_anti_rollback=True)
        changed=self.budget(data=budget_workbook(),dependency=dep)
        with self.assertRaisesRegex(service.FinanceRunError,'ANTI_ROLLBACK'): approve(changed,self.user)
        self.assertEqual(approve(changed,self.user,override_anti_rollback=True).status,'approved')

    def test_odd_cent_half_shares_replay_without_assigning_residuals_to_inputs(self):
        from openpyxl import load_workbook
        from api.tests_finance_upload_safety import workbook_bytes
        from masi_finance.publish.run_artifact import build_run_artifact
        workbook=load_workbook(BytesIO(workbook_bytes()))
        worksheet=workbook['Expenditure']
        worksheet['D2']=0.01;worksheet['J2']=0.01;worksheet['I2']='SYNTHETIC'
        output=BytesIO();workbook.save(output);workbook.close()
        artifact=build_run_artifact(output.getvalue(),source_name='20260907 - Ledger.xlsx',accounting_year=2026)
        dependency=approve(candidate(self.user,artifact=artifact),self.user,override_anti_rollback=True)
        run=self.budget(dependency=dependency)
        self.assertEqual([line['actual'] for line in run.payload['lines']],['0.01','0.01'])
        self.assertEqual([line['projected'] for line in run.payload['lines']],['0.02','0.02'])
        self.assertEqual(run.payload['lines_by_bc'][0]['monetary_projections']['actual']['residual'],'-0.01')
        with patch('openpyxl.load_workbook',side_effect=AssertionError('workbook is not retained')):
            self.assertEqual(approve(run,self.user).status,'approved')


class BudgetInstalledContractTests(TestCase):
    def test_schema_and_golden_are_exact_installed_contracts(self):
        from pathlib import Path
        from importlib.resources import files
        import json
        from masi_finance.publish.budget_run_schema import validate_budget_run_schema
        base=Path(__file__).parent
        resources=files('masi_finance.publish').joinpath('schema')
        for source,target in [('budget-run-1.0.0.json','contracts/budget-run-1.0.0.json'),
                              ('fixture-budget-run-1.0.0.json','tests_data/budget-run-1.0.0.json')]:
            self.assertEqual((base/target).read_bytes(),resources.joinpath(source).read_bytes())
        validate_budget_run_schema(json.loads((base/'tests_data/budget-run-1.0.0.json').read_bytes()))


from django.test import SimpleTestCase


class BudgetMigrationTests(SimpleTestCase):
    def test_upgrade_preserves_every_funder_field_and_constraint(self):
        # Isolate migration data from the main suite: TransactionTestCase flush
        # removes migration-seeded grants that the released fresh-install test
        # intentionally inspects. No shared database is mutated by this probe.
        import os
        import subprocess
        import sys
        script = """
import os, django, unittest
os.environ['DJANGO_SETTINGS_MODULE'] = 'masi_website.settings'
django.setup()
from django.core.management import call_command
call_command('migrate', verbosity=0)
from api.tests_finance_budgets import check_budget_migration
check_budget_migration(unittest.TestCase())
print('SQLite 0051 -> 0052: funder rows, JSON and constraints preserved')
"""
        result = subprocess.run([sys.executable, '-c', script],
                                env=dict(os.environ, DATABASE_URL='sqlite:///:memory:'),
                                capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('funder rows, JSON and constraints preserved', result.stdout)


def check_budget_migration(self):
    from django.db import connection, transaction, IntegrityError
    from django.db.migrations.executor import MigrationExecutor
    from api.finance_run_test_utils import legacy
    user=actor(); run=approve(candidate(user),user)
    row=legacy(2025)
    service.import_legacy_snapshots(user,year=2025,legacy_row_id=row.pk)
    executor=MigrationExecutor(connection)
    try:
        executor.migrate([('api','0051_finance_runs_foundation')])
        before_apps=executor.loader.project_state([('api','0051_finance_runs_foundation')]).apps
        before=before_apps.get_model('api','FinanceRun')
        retained=list(before.objects.order_by('pk').values())
        executor=MigrationExecutor(connection)
        executor.migrate([('api','0052_finance_budget_runs')])
        after=list(FinanceRun.objects.order_by('pk').values())
        self.assertEqual([{k:v for k,v in record.items() if k!='dependency_run_id'} for record in after],retained)
        self.assertTrue(all(record['dependency_run_id'] is None for record in after))
        for invalid in ({'kind':'unknown'}, {'payload':None}, {'approved_by':None}, {'status':'candidate'}):
            with self.assertRaises(IntegrityError), transaction.atomic():
                FinanceRun.objects.filter(pk=run.pk).update(**invalid)
        with self.assertRaises(IntegrityError), transaction.atomic():
            candidate(user)
    finally:
        executor=MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())
