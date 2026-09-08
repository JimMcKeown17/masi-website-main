"""Run history, checked publication and coherent current finance metadata."""
import json

from django.shortcuts import get_object_or_404
from rest_framework import serializers
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.pagination import CursorPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from api.models import FinanceRun
from api.permissions import IsFinancePublisher, IsFinanceReader, finance_capabilities_for
from api.services.finance_runs import (FinanceRunError, approve_run, demote_run, upload_workbook, pull_budget,
    validate_stored_run, budget_dependency_metadata)
from api.views.finance import AUTH_CLASSES


class RunPagination(CursorPagination):
    page_size = 50
    ordering = ('-uploaded_at', '-id')


def visible_runs(user):
    capabilities = finance_capabilities_for(user)
    if not capabilities:
        raise PermissionDenied('Finance access is not granted for this account.')
    statuses = []
    if 'finance.read' in capabilities:
        statuses.extend(['approved', 'superseded'])
    if 'finance.publish' in capabilities:
        statuses.extend(['candidate', 'failed'])
    return FinanceRun.objects.filter(status__in=statuses)


class RunMetadataSerializer(serializers.ModelSerializer):
    class Meta:
        model = FinanceRun
        fields = ('id', 'kind', 'accounting_year', 'status', 'source_name', 'source_date', 'source_sha256',
                  'source_size_bytes', 'schema_version', 'producer_version', 'payload_sha256', 'facts_sha256',
                  'uploaded_by', 'uploaded_at', 'approved_by', 'approved_at', 'previous_approved',
                  'approval_overrode_rollback', 'approval_acknowledged_findings', 'approval_note',
                  'demoted_by', 'demoted_at', 'demotion_note', 'parse_duration_ms', 'total_duration_ms',
                  'peak_memory_bytes', 'dependency_run', 'fact_row_count', 'allocation_count', 'finding_count', 'in_scope_error_count')
        read_only_fields = fields


def run_detail(run, user):
    result = dict(RunMetadataSerializer(run).data)
    result.update(manifest=run.manifest, payload=run.payload, failure=run.failure, allowed_actions=[])
    if 'finance.publish' in finance_capabilities_for(user):
        if run.status == 'candidate' or (run.status == 'superseded' and run.approved_at and run.approved_by_id):
            result['allowed_actions'].append('approve')
        if run.status == 'approved' and run.previous_approved_id:
            result['allowed_actions'].append('demote')
    return result


def year_parameter(params, *, required=False):
    value = params.get('year')
    if value is None and not required:
        return None
    try:
        year = int(value)
    except (ValueError, TypeError):
        raise ValidationError({'year': 'year must be an integer.'}) from None
    if not 1 <= year <= 32767:
        raise ValidationError({'year': 'year is outside the supported range.'})
    return year


class UploadMetadataSerializer(serializers.Serializer):
    kind = serializers.ChoiceField(choices=['funders', 'budgets'])
    year = serializers.IntegerField(min_value=2000, max_value=2100)
    source_name = serializers.CharField(max_length=255, trim_whitespace=False)
    client_modified_at = serializers.CharField(required=False, trim_whitespace=False)
    ledger_run_id = serializers.UUIDField(required=False)

    def to_internal_value(self, data):
        if set(data) - set(self.fields) or any(len(data.getlist(key)) != 1 for key in data):
            raise ValidationError('UPLOAD_METADATA_INVALID')
        return super().to_internal_value(data)


class BudgetPullSerializer(serializers.Serializer):
    year = serializers.IntegerField(min_value=2000, max_value=2100)
    ledger_run_id = serializers.UUIDField()


def _unique_json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError()
        result[key] = value
    return result


class FinanceBudgetPull(APIView):
    authentication_classes = AUTH_CLASSES
    permission_classes = [IsFinancePublisher]
    parser_classes = []

    def post(self, request):
        try:
            if request.content_type.split(';', 1)[0] != 'application/json':
                raise ValueError()
            raw = request._request
            stream = raw.environ['wsgi.input'] if hasattr(raw, 'environ') else raw
            body = bytearray()
            while len(body) <= 4096:
                count = 4097 - len(body)
                if hasattr(stream, '__len__'):
                    count = min(count, len(stream))
                chunk = stream.read(count)
                if not chunk:
                    break
                body.extend(chunk)
            if len(body) > 4096:
                raise ValueError()
            data = json.loads(body, object_pairs_hook=_unique_json_object)
        except (ValueError, TypeError, RecursionError):
            return Response({'code': 'UPLOAD_METADATA_INVALID'}, status=400)
        if (request.query_params or not isinstance(data, dict)
                or set(data) != {'year', 'ledger_run_id'} or type(data.get('year')) is not int):
            return Response({'code': 'UPLOAD_METADATA_INVALID'}, status=400)
        serializer = BudgetPullSerializer(data=data)
        if not serializer.is_valid():
            return Response({'code': 'UPLOAD_METADATA_INVALID'}, status=400)
        try:
            run, status = pull_budget(request.user, **serializer.validated_data)
        except FinanceRunError as error:
            return Response({'code': error.code}, status=error.status)
        return Response(run_detail(run, request.user), status=status)


class FinanceRunList(APIView):
    authentication_classes = AUTH_CLASSES
    # No DRF body parser: POST uses the underlying request stream directly.
    parser_classes = []

    def get_permissions(self):
        if self.request.method == 'POST':
            return [IsFinancePublisher()]
        return super().get_permissions()

    def post(self, request):
        serializer = UploadMetadataSerializer(data=request.query_params)
        if not serializer.is_valid():
            return Response({'code': 'UPLOAD_METADATA_INVALID'}, status=400)
        try:
            # DRF Request.stream consults Content-Length and may access body.
            # WSGIRequest's LimitedStream also trusts Content-Length. Use the
            # server-framed input, so missing/false lengths cannot truncate our
            # actual-byte counter. The configured deployment is WSGI.
            raw = request._request
            stream = raw.environ['wsgi.input'] if hasattr(raw, 'environ') else raw
            run, status = upload_workbook(
                stream, request.user, **serializer.validated_data,
                content_type=request.META.get('CONTENT_TYPE', ''),
                content_length=request.META.get('CONTENT_LENGTH'))
        except FinanceRunError as error:
            return Response({'code': error.code}, status=error.status)
        return Response(run_detail(run, request.user), status=status)

    def get(self, request):
        queryset = visible_runs(request.user)
        unknown = set(request.query_params) - {'kind', 'year', 'status', 'cursor'}
        if unknown:
            raise ValidationError('Unknown run filter.')
        if 'kind' in request.query_params:
            if request.query_params['kind'] not in ('funders', 'budgets'):
                raise ValidationError({'kind': 'Unsupported finance kind.'})
            queryset = queryset.filter(kind=request.query_params['kind'])
        year = year_parameter(request.query_params)
        if year is not None:
            queryset = queryset.filter(accounting_year=year)
        if 'status' in request.query_params:
            status = request.query_params['status']
            if status not in ('candidate', 'approved', 'superseded', 'failed'):
                raise ValidationError({'status': 'Unsupported run status.'})
            queryset = queryset.filter(status=status)
        pagination = RunPagination()
        page = pagination.paginate_queryset(queryset.defer('manifest', 'payload', 'failure'), request, view=self)
        return pagination.get_paginated_response(RunMetadataSerializer(page, many=True).data)


class FinanceRunDetail(APIView):
    authentication_classes = AUTH_CLASSES

    def get(self, request, run_id):
        run = get_object_or_404(visible_runs(request.user), pk=run_id)
        if run.kind == 'budgets' and run.status != 'failed':
            try:
                validate_stored_run(run)
                authorize_dependency(run, request.user)
            except FinanceRunError as error:
                return Response({'code': error.code}, status=error.status)
        return Response(run_detail(run, request.user))


class ApprovalOptionsSerializer(serializers.Serializer):
    override_anti_rollback = serializers.BooleanField(default=False)
    acknowledge_findings = serializers.BooleanField(default=False)
    note = serializers.CharField(default='', allow_blank=True, max_length=10000)

    def to_internal_value(self, data):
        if not isinstance(data, dict) or set(data) - set(self.fields):
            raise ValidationError({'non_field_errors': ['Only override_anti_rollback, acknowledge_findings and note are accepted.']})
        for field in ('override_anti_rollback', 'acknowledge_findings'):
            if field in data and type(data[field]) is not bool:
                raise ValidationError({field: 'Must be a JSON boolean.'})
        if 'note' in data and not isinstance(data['note'], str):
            raise ValidationError({'note': 'Must be a string.'})
        return super().to_internal_value(data)


class FinanceRunApprove(APIView):
    authentication_classes = AUTH_CLASSES
    permission_classes = [IsFinancePublisher]
    transition = staticmethod(approve_run)

    def post(self, request, run_id):
        serializer = ApprovalOptionsSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        try:
            run = self.transition(run_id, request.user, **serializer.validated_data)
        except FinanceRunError as exc:
            return Response({'code': exc.code, 'detail': exc.code}, status=exc.status)
        return Response(run_detail(run, request.user))


class FinanceRunDemote(FinanceRunApprove):
    transition = staticmethod(demote_run)


def compatibility_result(runs):
    if not runs:
        return False, {'code': 'NO_APPROVED_RUNS', 'runs': {}}
    if any(not run.get('management_accounts_sha256') for run in runs.values()):
        return False, {'code': 'DEPENDENCY_UNRESOLVED', 'runs': runs}
    if len({run['management_accounts_sha256'] for run in runs.values()}) != 1:
        return False, {'code': 'SOURCE_MISMATCH', 'runs': runs}
    return True, None


def management_accounts_sha(run):
    if run.kind == 'funders':
        return run.source_sha256
    if run.kind == 'budgets':
        dependency = run.dependency_run
        if (dependency is None or dependency.kind != 'funders'
                or dependency.accounting_year != run.accounting_year
                or dependency.status not in ('approved', 'superseded')
                or dependency.schema_version != '2.0.0' or not dependency.facts_sha256
                or run.manifest.get('dependencies') != [budget_dependency_metadata(dependency)]):
            return None
        return dependency.source_sha256
    # Future kinds must declare their source explicitly. Ambiguity fails closed.
    dependencies = run.manifest.get('dependencies', [])
    sources = {item.get('source_sha256') for item in dependencies if item.get('kind') == 'funders'}
    return next(iter(sources)) if len(sources) == 1 else None


class FinanceCurrent(APIView):
    authentication_classes = AUTH_CLASSES
    permission_classes = [IsFinanceReader]

    def get(self, request):
        if set(request.query_params) - {'year'}:
            raise ValidationError('Unknown current-run filter.')
        year = year_parameter(request.query_params, required=True)
        runs = {}
        for run in FinanceRun.objects.filter(accounting_year=year, status='approved').defer('payload', 'failure'):
            runs[run.kind] = {'id': str(run.pk), 'source_sha256': run.source_sha256,
                              'management_accounts_sha256': management_accounts_sha(run),
                              'schema_version': run.schema_version, 'approved_at': run.approved_at.isoformat()}
            if run.kind == 'budgets':
                runs[run.kind].update(budget_source_sha256=run.source_sha256,
                                      dependency_run_id=str(run.dependency_run_id))
        compatible, reason = compatibility_result(runs)
        return Response({'accounting_year': year, 'runs': runs, 'compatible': compatible, 'compatibility_reason': reason})



def authorize_dependency(run, user):
    """Recheck both capabilities and pinned dependency on every budget read."""
    capabilities = finance_capabilities_for(user)
    required = 'finance.publish' if run.status in ('candidate', 'failed') else 'finance.read'
    if required not in capabilities:
        raise PermissionDenied('Finance access is not granted for this account.')
    dependency = run.dependency_run
    if (dependency is None or dependency.kind != 'funders'
            or dependency.status not in ('approved', 'superseded')
            or dependency.accounting_year != run.accounting_year):
        raise ValidationError('BUDGET_DEPENDENCY_INVALID')
    return dependency


class RowPagination(CursorPagination):
    page_size = 100
    ordering = ('date', 'sheet_row', 'row_key')
    offset_cutoff = 50000


class FinanceRunRows(APIView):
    authentication_classes = AUTH_CLASSES

    def filtered_rows(self, request, run_id):
        from masi_finance.publish.excel import excel_equal
        allowed = {'year', 'bc', 'cursor', 'format'}
        if (set(request.query_params) - allowed
                or any(len(request.query_params.getlist(k)) != 1 for k in request.query_params)):
            raise ValidationError('ROW_FILTER_INVALID')
        year = year_parameter(request.query_params, required=True)
        bc = request.query_params.get('bc')
        if bc is None or len(bc) > 256:
            raise ValidationError('ROW_FILTER_INVALID')
        run = get_object_or_404(visible_runs(request.user), pk=run_id)
        ledger = authorize_dependency(run, request.user) if run.kind == 'budgets' else run
        if ledger.schema_version != '2.0.0' or not ledger.facts_sha256 or ledger.status == 'failed':
            raise ValidationError('FACTS_UNAVAILABLE')
        if run.kind == 'budgets' and year != run.accounting_year:
            raise ValidationError('ROW_FILTER_INVALID')
        try:
            validate_stored_run(run)
        except FinanceRunError:
            raise ValidationError('RUN_INTEGRITY_INVALID') from None
        rows = ledger.ledger_rows.filter(year=year)
        # SQL equality would lose R32 numeric-text and case equivalence. Resolve
        # only the finite BC vocabulary, then let indexed SQL bound the rows.
        variants = [value for value in rows.values_list('bc', flat=True).distinct()
                    if value is not None and excel_equal(value, bc)]
        return run, ledger, rows.filter(bc__in=variants).order_by('date', 'sheet_row', 'row_key')

    def get(self, request, run_id):
        run, ledger, rows = self.filtered_rows(request, run_id)
        pagination = RowPagination()
        page = pagination.paginate_queryset(rows, request, view=self)
        result = pagination.get_paginated_response([ledger_row_document(row) for row in page])
        result.data.update(run_id=str(run.pk), ledger_run_id=str(ledger.pk),
                           management_accounts_sha256=ledger.source_sha256,
                           contributor_basis='full_ledger_amount_before_budget_share')
        return result


def ledger_row_document(row):
    from api.services.finance_runs import ROW_FIELDS, _money_string
    result = {field: getattr(row, field) for field in ROW_FIELDS}
    result['date'] = row.date.isoformat()
    result['amount'] = _money_string(row.amount)
    result['coverage_amount'] = _money_string(row.coverage_amount)
    return result


class FinanceRunRowsExport(FinanceRunRows):
    def get_content_negotiator(self):
        from rest_framework.negotiation import DefaultContentNegotiation

        class ExportNegotiation(DefaultContentNegotiation):
            def filter_renderers(self, renderers, format):
                # format belongs to the download contract, not DRF's renderer.
                return renderers

        return ExportNegotiation()

    def get(self, request, run_id):
        import csv
        from io import BytesIO, StringIO
        from django.http import HttpResponse
        from api.services.finance_runs import ROW_FIELDS
        run, ledger, rows = self.filtered_rows(request, run_id)
        format_name = request.query_params.get('format', 'csv')
        if format_name not in ('csv', 'xlsx') or 'cursor' in request.query_params:
            raise ValidationError('ROW_FILTER_INVALID')
        if rows.count() > 50000:
            raise ValidationError('ROW_EXPORT_LIMIT')
        # Escape formula-leading text in CSV; XLSX stores text explicitly.
        if format_name == 'csv':
            output = StringIO(); writer = csv.writer(output); writer.writerow(ROW_FIELDS)
            for row in rows.iterator(chunk_size=1000):
                doc = ledger_row_document(row)
                writer.writerow([("'" + doc[k] if isinstance(doc[k], str) and doc[k].startswith(('=', '+', '-', '@'))
                                  and k not in ('amount', 'coverage_amount') else doc[k]) for k in ROW_FIELDS])
            response = HttpResponse(output.getvalue(), content_type='text/csv')
        else:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            wb = Workbook(); sheet = wb.active; sheet.title = 'Rows'
            sheet.append(list(ROW_FIELDS))
            for row in rows.iterator(chunk_size=1000):
                doc = ledger_row_document(row); cells = []
                for key in ROW_FIELDS:
                    cell = WriteOnlyCell(sheet, value=doc[key])
                    if isinstance(doc[key], str): cell.data_type = 's'
                    cells.append(cell)
                sheet.append(cells)
            output = BytesIO(); wb.save(output); wb.close()
            response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="finance-{run.pk}-rows.{format_name}"'
        return response
